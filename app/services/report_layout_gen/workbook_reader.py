# -*- coding: utf-8 -*-
"""
Excel workbook reader using openpyxl.
Extracts sheet structure, cell styles, merged cells, dimensions, and page setup into SheetModel.
"""
import io
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Tuple, Union
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
from openpyxl.utils import get_column_letter

from .models import Cell, MergeRange, SheetModel, col_letter_to_idx, parse_range_string
from .units import excel_width_to_px, pt_to_px

OFFICE_THEME_COLORS = [
    "#FFFFFF", "#000000", "#E7E6E6", "#44546A", "#4472C4",
    "#ED7D31", "#A5A5A5", "#FFC000", "#5B9BD5", "#70AD47"
]


def extract_color_hex(color: Any) -> Optional[str]:
    """Convert openpyxl Color object into #RRGGBB hex string."""
    if not color:
        return None
    try:
        if getattr(color, "type", None) == "rgb" or hasattr(color, "rgb"):
            rgb = color.rgb
            if isinstance(rgb, str) and rgb:
                # Openpyxl ARGB is 8 characters (AARRGGBB), strip alpha
                if len(rgb) == 8:
                    return f"#{rgb[2:].upper()}"
                elif len(rgb) == 6:
                    return f"#{rgb.upper()}"
        if getattr(color, "theme", None) is not None:
            t_idx = int(color.theme)
            if 0 <= t_idx < len(OFFICE_THEME_COLORS):
                base_hex = OFFICE_THEME_COLORS[t_idx]
                tint = getattr(color, "tint", 0.0) or 0.0
                if tint and abs(tint) > 0.01:
                    r = int(base_hex[1:3], 16)
                    g = int(base_hex[3:5], 16)
                    b = int(base_hex[5:7], 16)
                    if tint > 0:
                        r = int(r + (255 - r) * tint)
                        g = int(g + (255 - g) * tint)
                        b = int(b + (255 - b) * tint)
                    else:
                        r = int(r * (1 + tint))
                        g = int(g * (1 + tint))
                        b = int(b * (1 + tint))
                    return f"#{min(255, max(0, r)):02X}{min(255, max(0, g)):02X}{min(255, max(0, b)):02X}"
                return base_hex
    except Exception:
        pass
    return None


def format_cell_value(val: Any, num_format: Optional[str] = None) -> Tuple[Optional[str], Any]:
    """Format cell value for display, returning (formatted_str, raw_val)."""
    if val is None:
        return None, None
    raw = val
    nf = str(num_format or "")
    if isinstance(val, (int, float)):
        if "%" in nf:
            return f"{val * 100:.2f}%", raw
        elif isinstance(val, int) or (isinstance(val, float) and val.is_integer()):
            if "," in nf or "#" in nf:
                return f"{int(val):,}", raw
            return str(int(val)), raw
        elif isinstance(val, float):
            return f"{val:,.2f}", raw
    elif hasattr(val, "strftime"):
        return val.strftime("%d/%m/%Y"), raw
    return str(val).strip(), raw


def list_sheets(file_or_path: Union[str, Path, bytes, BinaryIO]) -> List[Dict[str, Any]]:
    """
    List basic metadata for all worksheets in a workbook.
    Uses ultra-fast direct zipfile inspection (~0.005s) with openpyxl fallback.
    Returns: [{'name': sheet_name, 'max_row': r, 'max_col': c, 'merge_count': m}]
    """
    try:
        if isinstance(file_or_path, bytes):
            z = zipfile.ZipFile(io.BytesIO(file_or_path), "r")
        elif isinstance(file_or_path, io.BytesIO):
            z = zipfile.ZipFile(file_or_path, "r")
        elif isinstance(file_or_path, (str, Path)):
            z = zipfile.ZipFile(file_or_path, "r")
        else:
            z = zipfile.ZipFile(file_or_path, "r")

        with z:
            wb_rels_xml = z.read("xl/_rels/workbook.xml.rels")
            wb_rels = ET.fromstring(wb_rels_xml)
            rel_map = {
                r.get("Id"): r.get("Target")
                for r in wb_rels.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
            }
            wb_root = ET.fromstring(z.read("xl/workbook.xml"))
            sheets = []
            for s in wb_root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
                name = s.get("name")
                r_id = s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                target = rel_map.get(r_id, "")
                if not target.startswith("xl/"):
                    target = "xl/" + target.lstrip("/")

                max_r, max_c = 1, 1
                merge_count = 0
                try:
                    with z.open(target) as f:
                        chunk = f.read(8192)
                        import re
                        m = re.search(rb'<dimension ref="([^"]+)"', chunk)
                        if m:
                            ref = m.group(1).decode("utf-8")
                            if ":" in ref:
                                _, _, max_r, max_c = parse_range_string(ref)
                            else:
                                max_r, max_c = 1, 1
                        mc = re.search(rb'<mergeCells count="([0-9]+)"', chunk)
                        if mc:
                            merge_count = int(mc.group(1).decode("utf-8"))
                        elif b"<mergeCell " in chunk:
                            merge_count = len(re.findall(rb"<mergeCell ", chunk))
                except Exception:
                    pass
                sheets.append({
                    "name": name,
                    "max_row": max_r,
                    "max_col": max_c,
                    "merge_count": merge_count,
                })
            if sheets:
                return sheets
    except Exception:
        pass

    # Fallback to openpyxl
    if isinstance(file_or_path, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_or_path), data_only=True, keep_links=False)
    else:
        wb = openpyxl.load_workbook(file_or_path, data_only=True, keep_links=False)

    results = []
    for name in wb.sheetnames:
        ws = wb[name]
        results.append({
            "name": name,
            "max_row": ws.max_row or 0,
            "max_col": ws.max_column or 0,
            "merge_count": len(ws.merged_cells.ranges) if ws.merged_cells else 0,
        })
    wb.close()
    return results


def extract_sheet_drawings(
    file_or_path: Union[str, Path, bytes, BinaryIO],
    sheet_name: str
) -> List[Dict[str, Any]]:
    """
    Extract text boxes, shapes and drawing annotations from worksheet drawing ML.
    Openpyxl ignores drawings/shapes, but report templates (such as Vietnamese tax/Bravo reports)
    frequently use Shapes/Textboxes for Title, Company Name, Tax ID, and Summary labels.
    """
    z = None
    try:
        if isinstance(file_or_path, bytes):
            z = zipfile.ZipFile(io.BytesIO(file_or_path), "r")
        elif isinstance(file_or_path, io.BytesIO):
            z = zipfile.ZipFile(file_or_path, "r")
        elif isinstance(file_or_path, (str, Path)):
            z = zipfile.ZipFile(file_or_path, "r")
        else:
            return []
    except Exception:
        return []

    try:
        ns_wb = {
            "wb": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        }
        ns_rel = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
        ns_draw = {
            "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
        }

        if "xl/workbook.xml" not in z.namelist() or "xl/_rels/workbook.xml.rels" not in z.namelist():
            return []

        wb_tree = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rel_map = {
            r.attrib["Id"]: r.attrib["Target"]
            for r in wb_rels.findall("rel:Relationship", ns_rel)
            if "Id" in r.attrib and "Target" in r.attrib
        }

        target_sheet_path = None
        for sh in wb_tree.findall(".//wb:sheet", ns_wb):
            name = sh.attrib.get("name")
            rid = sh.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if name == sheet_name and rid in rel_map:
                t = rel_map[rid]
                target_sheet_path = t if t.startswith("xl/") else f"xl/{t.lstrip('/')}"
                break

        if not target_sheet_path:
            return []

        sheet_rel_path = f"xl/worksheets/_rels/{Path(target_sheet_path).name}.rels"
        if sheet_rel_path not in z.namelist():
            return []

        sheet_rels = ET.fromstring(z.read(sheet_rel_path))
        drawing_path = None
        for r in sheet_rels.findall("rel:Relationship", ns_rel):
            if "relationships/drawing" in r.attrib.get("Type", ""):
                dt = r.attrib.get("Target", "")
                drawing_path = f"xl/{dt.replace('../', '')}"
                break

        if not drawing_path or drawing_path not in z.namelist():
            return []

        root = ET.fromstring(z.read(drawing_path))
        drawings = []

        anchors = root.findall(".//xdr:twoCellAnchor", ns_draw) + root.findall(".//xdr:oneCellAnchor", ns_draw)
        for anc in anchors:
            from_r = int(anc.findtext("xdr:from/xdr:row", "0", ns_draw)) + 1
            from_c = int(anc.findtext("xdr:from/xdr:col", "0", ns_draw)) + 1
            to_r = int(anc.findtext("xdr:to/xdr:row", str(from_r - 1), ns_draw)) + 1
            to_c = int(anc.findtext("xdr:to/xdr:col", str(from_c - 1), ns_draw)) + 1

            paragraphs = anc.findall(".//a:p", ns_draw)
            p_texts = []
            algn = None
            bold = False
            italic = False
            sz = None
            fc = None

            for p in paragraphs:
                pPr = p.find("a:pPr", ns_draw)
                if pPr is not None and not algn:
                    a_attr = pPr.attrib.get("algn")
                    if a_attr == "ctr": algn = "center"
                    elif a_attr == "r": algn = "right"
                    elif a_attr == "l": algn = "left"

                runs = p.findall("a:r", ns_draw)
                r_texts = []
                for r in runs:
                    rPr = r.find("a:rPr", ns_draw)
                    if rPr is not None:
                        if rPr.attrib.get("b") in ("1", "true", "True"): bold = True
                        if rPr.attrib.get("i") in ("1", "true", "True"): italic = True
                        if rPr.attrib.get("sz") and not sz:
                            try: sz = round(float(rPr.attrib.get("sz")) / 100.0, 1)
                            except Exception: pass
                        srgb = rPr.find(".//a:srgbClr", ns_draw)
                        if srgb is not None and not fc:
                            fc = f"#{srgb.attrib.get('val')}"
                    t_el = r.find("a:t", ns_draw)
                    if t_el is not None and t_el.text:
                        r_texts.append(t_el.text)
                if r_texts:
                    p_texts.append("".join(r_texts))

            full_txt = "\n".join(p_texts).strip()
            if full_txt:
                drawings.append({
                    "from_row": from_r,
                    "from_col": from_c,
                    "to_row": to_r,
                    "to_col": to_c,
                    "text": full_txt,
                    "bold": bold,
                    "italic": italic,
                    "font_size": sz,
                    "align": algn or "left",
                    "color": fc
                })

        return drawings
    except Exception:
        return []
    finally:
        if z:
            try:
                z.close()
            except Exception:
                pass


def read_sheet(
    file_or_path: Union[str, Path, bytes, BinaryIO],
    sheet_name: Optional[str] = None
) -> SheetModel:
    """
    Read a specific worksheet into a SheetModel with rich formatting and metadata.
    If sheet_name is None, loads the active sheet.
    """
    if isinstance(file_or_path, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_or_path), data_only=True, keep_links=False)
    else:
        wb = openpyxl.load_workbook(file_or_path, data_only=True, keep_links=False)

    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    actual_name = ws.title
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1

    # 1. Parse merged ranges and build fast O(1) lookup map
    merges: List[MergeRange] = []
    merge_map: Dict[Tuple[int, int], MergeRange] = {}
    if ws.merged_cells:
        for rng in ws.merged_cells.ranges:
            m = MergeRange(
                min_row=rng.min_row,
                min_col=rng.min_col,
                max_row=rng.max_row,
                max_col=rng.max_col,
            )
            merges.append(m)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merge_map[(r, c)] = m

    # 2. Extract column widths in px
    default_col_w = ws.sheet_format.defaultColWidth if ws.sheet_format else None
    col_widths_px: Dict[int, int] = {}

    for dim in ws.column_dimensions.values():
        if dim.width is not None and dim.min is not None and dim.max is not None:
            px = excel_width_to_px(dim.width, default_col_w or 8.43)
            for c in range(dim.min, dim.max + 1):
                col_widths_px[c] = px

    for c in range(1, max_col + 1):
        if c not in col_widths_px:
            col_letter = get_column_letter(c)
            dim = ws.column_dimensions.get(col_letter)
            w = dim.width if dim and dim.width is not None else default_col_w
            col_widths_px[c] = excel_width_to_px(w, 8.43)

    # 3. Extract row heights in px (fast lookup)
    default_row_h = ws.sheet_format.defaultRowHeight if ws.sheet_format else None
    row_heights_px: Dict[int, int] = {}
    for r, dim in ws.row_dimensions.items():
        if getattr(dim, "height", None) is not None:
            row_heights_px[r] = pt_to_px(dim.height, 15.0)

    # 4. Hidden rows, columns, freeze panes & grid lines
    hidden_rows = [r for r, dim in ws.row_dimensions.items() if getattr(dim, "hidden", False)]
    hidden_cols_set = set()
    for k, dim in ws.column_dimensions.items():
        if getattr(dim, "hidden", False):
            try:
                if dim.min and dim.max:
                    for col_i in range(dim.min, dim.max + 1):
                        hidden_cols_set.add(col_i)
                else:
                    hidden_cols_set.add(col_letter_to_idx(str(k)))
            except Exception:
                pass
    hidden_cols = sorted(list(hidden_cols_set))

    freeze_panes = str(ws.freeze_panes) if ws.freeze_panes else None
    show_grid_lines = True
    if hasattr(ws, "sheet_view") and ws.sheet_view and getattr(ws.sheet_view, "showGridLines", None) is not None:
        show_grid_lines = bool(ws.sheet_view.showGridLines)

    # 5. Extract styles and cells
    styles: List[Dict[str, Any]] = []
    style_map: Dict[str, int] = {}

    def get_style_id(c_obj: Any) -> Optional[int]:
        st = {}
        f = c_obj.font
        if f:
            if f.bold: st["b"] = True
            if f.italic: st["i"] = True
            if f.underline: st["u"] = True
            if f.size: st["fs"] = round(float(f.size), 1)
            fc = extract_color_hex(f.color)
            if fc and fc not in ("#000000", "#000"): st["fc"] = fc

        fl = c_obj.fill
        if fl and getattr(fl, "fill_type", None) == "solid":
            bg = extract_color_hex(getattr(fl, "fgColor", None))
            if bg and bg not in ("#FFFFFF", "#FFF"): st["bg"] = bg

        al = c_obj.alignment
        if al:
            if al.horizontal and al.horizontal != "general": st["ha"] = al.horizontal
            if al.vertical and al.vertical not in ("bottom", "center"): st["va"] = al.vertical
            if al.wrap_text: st["wrap"] = True

        bdr = c_obj.border
        if bdr:
            bd = {}
            for side in ("top", "bottom", "left", "right"):
                s_obj = getattr(bdr, side, None)
                if s_obj and s_obj.style:
                    col = extract_color_hex(s_obj.color) or "#D4D4D8"
                    bd[side] = {"s": s_obj.style, "c": col}
            if bd:
                st["bd"] = bd

        if not st:
            return None

        key = str(sorted(st.items()))
        if key in style_map:
            return style_map[key]
        idx = len(styles)
        style_map[key] = idx
        styles.append(st)
        return idx

    cells: Dict[Tuple[int, int], Cell] = {}
    row_ranges = [(1, max_row)]
    if max_row > 500:
        row_ranges = [(1, min(max_row, 300)), (max(301, max_row - 60), max_row)]

    for start_r, end_r in row_ranges:
        for row in ws.iter_rows(min_row=start_r, max_row=end_r, min_col=1, max_col=max_col):
            for c in row:
                r_idx, c_idx = c.row, c.column
                merge_range = merge_map.get((r_idx, c_idx))
                # Skip blank unmerged unstyled cells to optimize loading speed and memory
                if c.value is None and not merge_range and not c.has_style:
                    continue

                formatted_val, raw_val = format_cell_value(c.value, getattr(c, "number_format", None))
                is_anchor = True
                if merge_range:
                    is_anchor = (r_idx == merge_range.min_row and c_idx == merge_range.min_col)

                sid = get_style_id(c)
                cells[(r_idx, c_idx)] = Cell(
                    row=r_idx,
                    col=c_idx,
                    value=formatted_val,
                    raw_value=raw_val,
                    formatted_value=formatted_val,
                    merge=merge_range,
                    is_anchor=is_anchor,
                    style_id=sid,
                )

    # 5b. Extract and integrate drawing/shape texts (floating text boxes, titles, tax codes)
    drawing_shapes = extract_sheet_drawings(file_or_path, actual_name)
    if drawing_shapes:
        # Group single digit shapes on the same row (e.g. tax code boxes)
        processed_shapes = []
        digit_shapes_by_row: Dict[int, List[Dict[str, Any]]] = {}
        for s in drawing_shapes:
            txt = s["text"].strip()
            if len(txt) == 1 and txt.isdigit():
                digit_shapes_by_row.setdefault(s["from_row"], []).append(s)
            else:
                processed_shapes.append(s)

        for r, d_shapes in digit_shapes_by_row.items():
            if len(d_shapes) > 1:
                # Combine tax code digits into a cleanly spaced string
                d_shapes.sort(key=lambda x: x["from_col"])
                combined_digits = "   ".join(d["text"].strip() for d in d_shapes)
                first_d = d_shapes[0]
                min_c = min(d["from_col"] for d in d_shapes)
                max_c = max(d["to_col"] for d in d_shapes)
                processed_shapes.append({
                    "from_row": r,
                    "from_col": min_c,
                    "to_row": r,
                    "to_col": max_c,
                    "text": combined_digits,
                    "bold": True,
                    "italic": False,
                    "font_size": first_d.get("font_size") or 11.0,
                    "align": "left",
                    "color": first_d.get("color"),
                })
            else:
                processed_shapes.extend(d_shapes)

        # Inject processed shapes into cells
        for shape in processed_shapes:
            r = shape["from_row"]
            c = shape["from_col"]
            to_r = shape["to_row"]
            to_c = shape["to_col"]

            # Only inject if target cell has no value
            existing_cell = cells.get((r, c))
            if existing_cell and existing_cell.value not in (None, ""):
                continue

            # Determine span
            span_r = max(r, to_r - 1) if to_r > r else r
            span_c = max(c, to_c - 1) if to_c > c else c

            max_row = max(max_row, span_r)
            max_col = max(max_col, span_c)

            # Create or lookup style for drawing text
            st = {}
            if shape.get("bold"): st["b"] = True
            if shape.get("italic"): st["i"] = True
            if shape.get("font_size"): st["fs"] = shape["font_size"]
            if shape.get("align"): st["ha"] = shape["align"]
            if shape.get("color") and shape["color"] not in ("#000000", "#000"): st["fc"] = shape["color"]

            shape_sid = None
            if st:
                st_key = str(sorted(st.items()))
                if st_key in style_map:
                    shape_sid = style_map[st_key]
                else:
                    shape_sid = len(styles)
                    style_map[st_key] = shape_sid
                    styles.append(st)

            # Create merge range if spans more than 1 cell and not already merged
            shape_merge = None
            if (span_r > r or span_c > c) and not merge_map.get((r, c)):
                shape_merge = MergeRange(min_row=r, min_col=c, max_row=span_r, max_col=span_c)
                merges.append(shape_merge)
                for mr in range(r, span_r + 1):
                    for mc in range(c, span_c + 1):
                        merge_map[(mr, mc)] = shape_merge

            # Set anchor cell
            txt = shape["text"]
            cells[(r, c)] = Cell(
                row=r,
                col=c,
                value=txt,
                raw_value=txt,
                formatted_value=txt,
                merge=shape_merge or merge_map.get((r, c)),
                is_anchor=True,
                style_id=shape_sid,
            )

            # Set covered cells if merged
            if shape_merge:
                for mr in range(r, span_r + 1):
                    for mc in range(c, span_c + 1):
                        if (mr, mc) == (r, c):
                            continue
                        cells[(mr, mc)] = Cell(
                            row=mr,
                            col=mc,
                            value=None,
                            raw_value=None,
                            formatted_value=None,
                            merge=shape_merge,
                            is_anchor=False,
                            style_id=shape_sid,
                        )

    # 6. Page setup details
    page_info: Dict[str, Any] = {
        "orientation": ws.page_setup.orientation if ws.page_setup else "portrait",
        "paper_size": ws.page_setup.paperSize if ws.page_setup else 9,  # 9 = A4, 11 = A5
        "margins": {
            "top": ws.page_margins.top if ws.page_margins else 0.75,
            "bottom": ws.page_margins.bottom if ws.page_margins else 0.75,
            "left": ws.page_margins.left if ws.page_margins else 0.7,
            "right": ws.page_margins.right if ws.page_margins else 0.7,
            "header": ws.page_margins.header if ws.page_margins else 0.3,
            "footer": ws.page_margins.footer if ws.page_margins else 0.3,
        } if ws.page_margins else {},
    }

    wb.close()

    default_col_px = excel_width_to_px(default_col_w or 8.43, 8.43)
    default_row_px = pt_to_px(default_row_h or 15.0, 15.0)

    return SheetModel(
        name=actual_name,
        max_row=max_row,
        max_col=max_col,
        cells=cells,
        merges=merges,
        col_widths_px=col_widths_px,
        row_heights_px=row_heights_px,
        styles=styles,
        hidden_rows=hidden_rows,
        hidden_cols=hidden_cols,
        freeze_panes=freeze_panes,
        default_col_width_px=default_col_px,
        default_row_height_px=default_row_px,
        show_grid_lines=show_grid_lines,
        page=page_info,
        drawings=drawing_shapes,
    )
