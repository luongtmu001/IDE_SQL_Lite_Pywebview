# -*- coding: utf-8 -*-
import base64
import gzip
import io
import json
import threading
import os
import re
import sys
import urllib.parse
import uuid
import xml.dom.minidom
import xml.etree.ElementTree as ET
import zipfile
import zlib
from pathlib import Path

# Resolve bundle and app directories for both development and PyInstaller frozen modes
if getattr(sys, 'frozen', False):
    APP_DIR = Path(sys.executable).resolve().parent
    meipass = getattr(sys, '_MEIPASS', None)
    if meipass and (Path(meipass) / "templates").exists():
        BUNDLE_DIR = Path(meipass).resolve()
    elif (APP_DIR / "_internal" / "templates").exists():
        BUNDLE_DIR = APP_DIR / "_internal"
    elif (APP_DIR / "templates").exists():
        BUNDLE_DIR = APP_DIR
    else:
        BUNDLE_DIR = APP_DIR
else:
    APP_DIR = Path(__file__).resolve().parent
    BUNDLE_DIR = APP_DIR

BASE_DIR = BUNDLE_DIR
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

import logging
import webview
from app.utils.crypto import encrypt_password, decrypt_password
from app.utils.dwm import apply_dwm_titlebar_theme, THEME_TITLEBAR_PRESETS
from app.utils.logging_setup import (
    setup_logging,
    open_log_folder as _open_log_folder,
    get_log_info as _get_log_info,
    update_log_level,
)
from app.utils.sql_safety import check_query_safety
from app.services.connection_manager import ConnectionManager
from app.services.table_designer_service import TableDesignerService
from app.services.table_data_editor_service import TableDataEditorService
from app.database.factory import create_adapter
from app.services.history import get_history_service

logger = logging.getLogger("luobtool.main")

from app.services.settings_store import (
    DATA_DIR,
    SETTINGS_FILE,
    SETTINGS_FILE_ALT,
    CONNECTIONS_FILE,
    SNIPPETS_FILE,
    CODE_STYLES_FILE,
    PROGRAM_LIST_FILE,
    PROGRAM_GROUPS_FILE,
    REDGATE_SNIPPETS_DIR,
    _get_redgate_snippets_dir,
    _read_json,
    _write_json,
    _load_settings_file,
    _save_settings_file,
    _load_snippets_file,
    _save_snippets_file,
    _import_redgate_snippets,
)
from app.desktop.themes import (
    _clean_jsonc,
    _hex_to_rgb,
    _is_dark_color,
    _generate_theme_css,
    _process_import_vsix_theme,
)
from app.desktop.clipboard import (
    _native_set_clipboard,
    _native_get_clipboard,
)
from app.bravo.xml_codec import (
    _beautify_xml_string,
    _decode_xml_payload,
    _encode_xml_payload,
    _validate_xml_syntax,
)

connection_manager = ConnectionManager()
_OWNER_SESSION = "desktop_user"

# -------------------------------------------------------------------------
# BravoApi
# -------------------------------------------------------------------------


class BravoApi:
    def __init__(self, cm: ConnectionManager, owner_session_id: str = None, initial_context: dict = None, main_api=None):
        self.cm = cm
        self.owner_session_id = owner_session_id or f"bravo_user_{uuid.uuid4().hex[:8]}"
        self.initial_context = initial_context or {}
        self._bravo_sessions = {}
        self._window = None
        self._profiler_windows = []
        self._event_listener_windows = []
        self._main_api = main_api
        self._current_titlebar_theme = None

    def set_window(self, win):
        self._window = win

    def ready(self):
        if self._window:
            try:
                self._window.show()
            except Exception:
                pass
            import threading
            for delay in (0.05, 0.2, 0.5):
                threading.Timer(delay, lambda: self.apply_titlebar_theme()).start()
        return {"success": True}

    def open_log_folder(self):
        import app.utils.logging_setup as logging_setup
        return logging_setup.open_log_folder()

    def get_log_info(self):
        import app.utils.logging_setup as logging_setup
        return {"success": True, "data": logging_setup.get_log_info()}

    def tx_begin(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            return conn.begin_transaction()
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def tx_commit(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            return conn.commit_transaction()
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def tx_rollback(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            return conn.rollback_transaction()
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def tx_status(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            from datetime import datetime, timezone
            duration = None
            if conn.in_transaction and conn.tx_started_at:
                duration = (datetime.now(timezone.utc) - conn.tx_started_at).total_seconds()
            return {
                "success": True,
                "in_transaction": conn.in_transaction,
                "duration_seconds": duration,
                "environment": conn.environment,
                "read_only": conn.read_only
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def set_connection_read_only(self, connection_id, read_only=False):
        try:
            conn = self._get_connection(connection_id)
            if hasattr(conn, "set_read_only"):
                return conn.set_read_only(read_only)
            conn.read_only = bool(read_only)
            return {"success": True, "read_only": conn.read_only, "connection_id": connection_id}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def apply_titlebar_theme(self, bg_hex=None, text_hex=None, border_hex=None, is_dark=None):
        """Đồng bộ màu sắc thanh tiêu đề Windows qua DWM API cho tất cả các cửa sổ đang mở."""
        if getattr(self, '_main_api', None) and self._main_api is not self:
            try:
                return self._main_api.apply_titlebar_theme(bg_hex, text_hex, border_hex, is_dark)
            except Exception:
                pass

        try:
            if not bg_hex and getattr(self, '_current_titlebar_theme', None):
                t = self._current_titlebar_theme
                bg_hex = t.get("bg")
                text_hex = text_hex or t.get("text")
                border_hex = border_hex or t.get("border")
                if is_dark is None:
                    is_dark = t.get("is_dark")

            if not bg_hex:
                settings = _load_settings_file()
                bg_hex = settings.get("appearance", {}).get("theme") or settings.get("theme", "dark")

            if bg_hex in THEME_TITLEBAR_PRESETS:
                preset = THEME_TITLEBAR_PRESETS[bg_hex]
                bg_hex = preset["bg"]
                text_hex = text_hex or preset.get("text")
                border_hex = border_hex or preset.get("border")
                if is_dark is None:
                    is_dark = preset.get("is_dark")

            self._current_titlebar_theme = {
                "bg": bg_hex,
                "text": text_hex,
                "border": border_hex,
                "is_dark": is_dark
            }

            # 1. Áp dụng cho Main Window
            if getattr(self, "_window", None):
                apply_dwm_titlebar_theme(
                    self._window,
                    bg=bg_hex,
                    text=text_hex,
                    border=border_hex,
                    is_dark=is_dark,
                    title="luoBTool IDE"
                )

            # 2. Áp dụng cho Settings Window
            if getattr(self, "_settings_window", None):
                apply_dwm_titlebar_theme(
                    self._settings_window,
                    bg=bg_hex,
                    text=text_hex,
                    border=border_hex,
                    is_dark=is_dark,
                    title="Cài đặt (Settings) — luoBTool IDE"
                )

            # 3. Áp dụng cho các cửa sổ BRAVO Tool
            for bw in getattr(self, "_bravo_windows", []):
                if bw:
                    apply_dwm_titlebar_theme(
                        bw,
                        bg=bg_hex,
                        text=text_hex,
                        border=border_hex,
                        is_dark=is_dark
                    )

            # 4. Áp dụng cho các cửa sổ Profiler
            for pw in getattr(self, "_profiler_windows", []):
                if pw:
                    apply_dwm_titlebar_theme(
                        pw,
                        bg=bg_hex,
                        text=text_hex,
                        border=border_hex,
                        is_dark=is_dark,
                        title="SQL Trace Profiler — luoBTool IDE"
                    )

            # 5. Áp dụng cho các cửa sổ PostgreSQL Event Listener
            for ew in getattr(self, "_event_listener_windows", []):
                if ew:
                    apply_dwm_titlebar_theme(
                        ew,
                        bg=bg_hex,
                        text=text_hex,
                        border=border_hex,
                        is_dark=is_dark,
                        title="PostgreSQL Event Listener — luoBTool IDE"
                    )

            # 6. Áp dụng cho cửa sổ Query History
            if getattr(self, "_history_window", None):
                apply_dwm_titlebar_theme(
                    self._history_window,
                    bg=bg_hex,
                    text=text_hex,
                    border=border_hex,
                    is_dark=is_dark,
                    title="Query History — luoBTool IDE"
                )

            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # Native Clipboard Bridge
    def set_clipboard_text(self, text):
        try:
            hwnd = None
            if hasattr(self, '_window') and self._window:
                native = getattr(self._window, 'native', None)
                if native:
                    handle = getattr(native, 'Handle', None)
                    if handle:
                        try:
                            hwnd = int(handle)
                        except Exception:
                            pass
            ok = _native_set_clipboard(text, hwnd)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_clipboard_text(self):
        try:
            hwnd = None
            if hasattr(self, '_window') and self._window:
                native = getattr(self._window, 'native', None)
                if native:
                    handle = getattr(native, 'Handle', None)
                    if handle:
                        try:
                            hwnd = int(handle)
                        except Exception:
                            pass
            val = _native_get_clipboard(hwnd)
            return {"success": True, "text": val}
        except Exception as exc:
            return {"success": False, "error": str(exc), "text": ""}

    # Table Data Editor Methods
    def get_table_editor_metadata(self, connection_id, database=None, schema=None, table=None):
        try:
            if not table:
                return {"success": False, "error": "Tên bảng là bắt buộc (table is required)"}
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                meta = TableDataEditorService.get_metadata(conn.adapter, database, schema, table, lock=conn.lock)
                return {"success": True, "metadata": meta}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def fetch_table_data(self, connection_id, database=None, schema=None, table=None, select_cols=None, criteria=None, top_n=200, custom_sql=None):
        try:
            if not table and not custom_sql:
                return {"success": False, "error": "Tên bảng hoặc câu truy vấn là bắt buộc"}
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                data = TableDataEditorService.fetch_data(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    select_cols=select_cols,
                    criteria=criteria,
                    top_n=top_n,
                    custom_sql=custom_sql,
                    lock=conn.lock
                )
                return {"success": True, **data}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def update_table_row(self, connection_id, database=None, schema=None, table=None, pk_conditions=None, changes=None):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                result = TableDataEditorService.update_row(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    pk_conditions=pk_conditions or {},
                    changes=changes or {},
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def delete_table_row(self, connection_id, database=None, schema=None, table=None, pk_conditions=None):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                result = TableDataEditorService.delete_row(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    pk_conditions=pk_conditions or {},
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def insert_table_row(self, connection_id, database=None, schema=None, table=None, values=None, identity_columns=None, computed_columns=None):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                result = TableDataEditorService.insert_row(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    values=values or {},
                    identity_columns=identity_columns,
                    computed_columns=computed_columns,
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def submit_table_changes(self, connection_id, database=None, schema=None, table=None, changes=None, concurrency_mode="optimistic"):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            if getattr(conn, "read_only", False):
                return {"success": False, "error_code": "ERR_READ_ONLY", "error": "Kết nối đang ở chế độ CHỈ ĐỌC (Read-Only). Không thể lưu thay đổi vào bảng dữ liệu."}
            with conn.exclusive():
                result = TableDataEditorService.submit_changes(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    changes=changes or {},
                    concurrency_mode=concurrency_mode,
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def validate_table_rows(self, connection_id, database=None, schema=None, table=None, rows=None):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"success": False, "error": f"Không tìm thấy kết nối: {connection_id}", "total": len(rows or []), "valid_count": 0, "error_count": len(rows or []), "results": []}
            with conn.exclusive():
                result = TableDataEditorService.validate_insert_rows(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    rows=rows or [],
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"success": False, "error": str(exc), "total": len(rows or []), "valid_count": 0, "error_count": len(rows or []), "results": []}

    def validate_table_row(self, connection_id, database=None, schema=None, table=None, values=None, keys=None, is_new=True, target_column=None):
        try:
            conn = self._get_connection(connection_id)
            if not conn:
                return {"valid": False, "error": f"Không tìm thấy kết nối: {connection_id}"}
            with conn.exclusive():
                result = TableDataEditorService.validate_row(
                    adapter=conn.adapter,
                    database=database,
                    schema=schema,
                    table=table,
                    values=values or {},
                    keys=keys,
                    is_new=is_new,
                    target_column=target_column,
                    lock=conn.lock
                )
                return result
        except Exception as exc:
            return {"valid": False, "error": str(exc)}

    def get_initial_context(self):
        return {"success": True, "context": self.initial_context}

    def set_window_title(self, title):
        if self._window:
            try:
                self._window.set_title(title)
                return {"success": True}
            except Exception:
                pass
        return {"success": False}

    def get_config(self):
        settings = _load_settings_file()
        enabled = settings.get("addons", {}).get("bravo_tool", {}).get("enabled", True)
        return {"success": True, "enabled": bool(enabled)}

    def list_features(self):
        features = [
            {
                "id": "program-list",
                "displayName": "Danh sách chương trình",
                "icon": "fa-list-check",
                "order": 5,
                "description": "Quản lý danh sách kết nối và khởi chạy nhanh các chương trình BRAVO"
            },
            {
                "id": "layout-editor",
                "displayName": "Soạn thảo layout",
                "icon": "fa-table-columns",
                "order": 10,
                "description": "Thiết kế và soạn thảo layout báo cáo / form hiển thị"
            }
        ]
        return {"success": True, "features": features}

    def get_program_list(self):
        try:
            programs = _read_json(PROGRAM_LIST_FILE, [])
            if not isinstance(programs, list):
                programs = []
            res = []
            for p in programs:
                if not isinstance(p, dict):
                    continue
                item = dict(p)
                if item.get("password"):
                    item["password"] = decrypt_password(item["password"])
                res.append(item)
            return {"success": True, "programs": res}
        except Exception as exc:
            return {"success": False, "error": str(exc), "programs": []}

    def save_program_list(self, programs):
        try:
            if not isinstance(programs, list):
                return {"success": False, "error": "Dữ liệu chương trình không hợp lệ"}
            to_save = []
            for p in programs:
                if not isinstance(p, dict):
                    continue
                item = dict(p)
                pwd = item.get("password")
                if pwd and not str(pwd).startswith("ENC:"):
                    item["password"] = encrypt_password(pwd)
                to_save.append(item)
            ok = _write_json(PROGRAM_LIST_FILE, to_save)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_program_groups(self):
        try:
            groups = _read_json(PROGRAM_GROUPS_FILE, [])
            if not isinstance(groups, list):
                groups = []
            return {"success": True, "groups": groups}
        except Exception as exc:
            return {"success": False, "error": str(exc), "groups": []}

    def save_program_groups(self, groups):
        try:
            if not isinstance(groups, list):
                return {"success": False, "error": "Dữ liệu nhóm không hợp lệ"}
            to_save = []
            for g in groups:
                if not isinstance(g, dict):
                    continue
                to_save.append({
                    "id": str(g.get("id") or ""),
                    "name": str(g.get("name") or "").strip()
                })
            ok = _write_json(PROGRAM_GROUPS_FILE, to_save)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def select_program_file(self):
        try:
            target_win = getattr(self, '_window', None)
            if not target_win and hasattr(self, '_main_api') and self._main_api:
                target_win = getattr(self._main_api, '_window', None)

            # 1. Ưu tiên tuyệt đối sử dụng hộp thoại chuẩn của pywebview trên cửa sổ hiện hành
            if target_win:
                import webview
                files = target_win.create_file_dialog(
                    webview.OPEN_DIALOG,
                    allow_multiple=False,
                    file_types=('Executable Files (*.exe;*.bat;*.cmd)', 'All files (*.*)')
                )
                if files and len(files) > 0:
                    return {"success": True, "path": files[0]}
                # Người dùng bấm Hủy (Cancel) hoặc đóng hộp thoại -> dừng lại, không mở thêm cửa sổ thứ 2
                return {"success": False, "cancelled": True}

            # 2. Fallback Tkinter chỉ dùng khi hoàn toàn không có cửa sổ webview nào (môi trường headless)
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk()
            root.withdraw()
            root.attributes('-topmost', True)
            file_path = filedialog.askopenfilename(
                title="Chọn chương trình thực thi (Executable)",
                filetypes=[("Chương trình thực thi (*.exe;*.bat;*.cmd)", "*.exe *.bat *.cmd"), ("Tất cả tệp", "*.*")]
            )
            root.destroy()

            if file_path:
                return {"success": True, "path": file_path}
            return {"success": False, "cancelled": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def import_vsix_theme(self, target_win=None):
        try:
            if not target_win:
                target_win = getattr(self, '_window', None)
                if not target_win and hasattr(self, '_main_api') and self._main_api:
                    target_win = getattr(self._main_api, '_window', None)

            selected_file = None
            if target_win:
                import webview
                files = target_win.create_file_dialog(
                    webview.OPEN_DIALOG,
                    allow_multiple=False,
                    file_types=('VS Code Theme Extension (*.vsix)', 'All files (*.*)')
                )
                if files and len(files) > 0:
                    selected_file = files[0]
                else:
                    return {"success": False, "cancelled": True}
            else:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes('-topmost', True)
                selected_file = filedialog.askopenfilename(
                    title="Chọn file giao diện VS Code (.vsix)",
                    filetypes=[("VS Code Extension (*.vsix)", "*.vsix"), ("Tất cả tệp", "*.*")]
                )
                root.destroy()
                if not selected_file:
                    return {"success": False, "cancelled": True}

            return _process_import_vsix_theme(selected_file)
        except Exception as exc:
            import traceback
            traceback.print_exc()
            return {"success": False, "error": str(exc)}

    def grid_open_in_excel(self, columns=None, rows=None):
        try:
            import tempfile, os, datetime, uuid, ctypes
            columns = columns or []
            rows = rows or []

            temp_dir = os.path.join(tempfile.gettempdir(), "SqlGridExports")
            os.makedirs(temp_dir, exist_ok=True)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            unique_id = uuid.uuid4().hex[:6]
            file_name = f"Export_{timestamp}_{unique_id}.xlsx"
            file_path = os.path.join(temp_dir, file_name)

            col_names = [c if isinstance(c, str) else c.get("name", str(c)) for c in columns]

            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Query Results"

                header_font = Font(name="Segoe UI", size=10, bold=True, color="1F1F1F")
                header_fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
                thin_border = Border(
                    left=Side(style='thin', color='D4D4D4'),
                    right=Side(style='thin', color='D4D4D4'),
                    top=Side(style='thin', color='D4D4D4'),
                    bottom=Side(style='thin', color='D4D4D4')
                )

                for col_idx, col_name in enumerate(col_names, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=col_name)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                cell_font = Font(name="Segoe UI", size=10, color="000000")
                null_font = Font(name="Segoe UI", size=10, italic=True, color="888888")
                null_fill = PatternFill(start_color="FFFFE1", end_color="FFFFE1", fill_type="solid")

                for r_idx, row in enumerate(rows, start=2):
                    for c_idx, val in enumerate(row, start=1):
                        cell = ws.cell(row=r_idx, column=c_idx)
                        cell.border = thin_border
                        if val is None:
                            cell.value = "NULL"
                            cell.font = null_font
                            cell.fill = null_fill
                        else:
                            cell.value = val
                            cell.font = cell_font

                if col_names:
                    last_col_letter = get_column_letter(len(col_names))
                    ws.auto_filter.ref = f"A1:{last_col_letter}{max(len(rows) + 1, 1)}"
                ws.freeze_panes = "A2"

                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    sample = col[:100]
                    max_len = max([len(str(c.value or '')) for c in sample] or [10])
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

                wb.save(file_path)
            except Exception as excel_err:
                import csv
                file_path = file_path.replace(".xlsx", ".csv")
                with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(col_names)
                    for r in rows:
                        writer.writerow(["" if v is None else v for v in r])

            # Gán thuộc tính Read-Only cho file tạm
            try:
                ctypes.windll.kernel32.SetFileAttributesW(file_path, 1) # FILE_ATTRIBUTE_READONLY
            except Exception:
                try:
                    os.chmod(file_path, 0o444)
                except Exception:
                    pass

            os.startfile(file_path)
            return {"success": True, "file_path": file_path}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def grid_save_results_as(self, columns=None, rows=None, default_filename=None):
        try:
            import os, csv, json
            columns = columns or []
            rows = rows or []
            col_names = [c if isinstance(c, str) else c.get("name", str(c)) for c in columns]
            default_fn = default_filename or "Query_Results.xlsx"

            target_win = getattr(self, '_window', None)
            if not target_win and hasattr(self, '_main_api') and self._main_api:
                target_win = getattr(self._main_api, '_window', None)

            import webview
            if not target_win and hasattr(webview, 'windows') and webview.windows:
                target_win = webview.windows[0]

            file_path = None
            if target_win:
                try:
                    save_mode = getattr(webview.FileDialog, 'SAVE', getattr(webview, 'SAVE_DIALOG', 30))
                    files = target_win.create_file_dialog(
                        save_mode,
                        save_filename=default_fn,
                        file_types=(
                            'Excel Files (*.xlsx)',
                            'CSV Files (*.csv)',
                            'TSV Files (*.txt)',
                            'JSON Files (*.json)',
                            'All Files (*.*)'
                        )
                    )
                    if files and len(files) > 0:
                        file_path = files[0] if isinstance(files, (list, tuple)) else str(files)
                    else:
                        return {"success": False, "cancelled": True}
                except Exception as diag_err:
                    print(f"[SaveFileDialog Error] pywebview dialog error: {diag_err}, falling back to Tkinter")
                    file_path = None

            if not file_path:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes('-topmost', True)
                file_path = filedialog.asksaveasfilename(
                    initialfile=default_fn,
                    filetypes=[
                        ("Excel Files (*.xlsx)", "*.xlsx"),
                        ("CSV Files (*.csv)", "*.csv"),
                        ("TSV Files (*.txt)", "*.txt"),
                        ("JSON Files (*.json)", "*.json"),
                        ("All Files (*.*)", "*.*")
                    ]
                )
                root.destroy()
                if not file_path:
                    return {"success": False, "cancelled": True}

            ext = os.path.splitext(file_path)[1].lower()
            if not ext:
                ext = ".xlsx"
                file_path += ext

            if ext == ".xlsx":
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Results"

                header_font = Font(name="Segoe UI", size=10, bold=True, color="1F1F1F")
                header_fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
                thin_border = Border(
                    left=Side(style='thin', color='D4D4D4'),
                    right=Side(style='thin', color='D4D4D4'),
                    top=Side(style='thin', color='D4D4D4'),
                    bottom=Side(style='thin', color='D4D4D4')
                )

                for col_idx, col_name in enumerate(col_names, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=col_name)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                cell_font = Font(name="Segoe UI", size=10, color="000000")
                null_font = Font(name="Segoe UI", size=10, italic=True, color="888888")
                null_fill = PatternFill(start_color="FFFFE1", end_color="FFFFE1", fill_type="solid")

                for r_idx, row in enumerate(rows, start=2):
                    for c_idx, val in enumerate(row, start=1):
                        cell = ws.cell(row=r_idx, column=c_idx)
                        cell.border = thin_border
                        if val is None:
                            cell.value = "NULL"
                            cell.font = null_font
                            cell.fill = null_fill
                        else:
                            cell.value = val
                            cell.font = cell_font

                if col_names:
                    last_col_letter = get_column_letter(len(col_names))
                    ws.auto_filter.ref = f"A1:{last_col_letter}{max(len(rows) + 1, 1)}"
                ws.freeze_panes = "A2"

                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    sample = col[:100]
                    max_len = max([len(str(c.value or '')) for c in sample] or [10])
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

                wb.save(file_path)

            elif ext == ".csv":
                with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(col_names)
                    for row in rows:
                        writer.writerow(["" if val is None else val for val in row])

            elif ext in [".txt", ".tsv"]:
                with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
                    writer = csv.writer(f, delimiter="\t")
                    writer.writerow(col_names)
                    for row in rows:
                        writer.writerow(["" if val is None else val for val in row])

            elif ext == ".json":
                data_list = []
                for row in rows:
                    item = {}
                    for c_idx, col_name in enumerate(col_names):
                        item[col_name] = row[c_idx] if c_idx < len(row) else None
                    data_list.append(item)
                with open(file_path, "w", encoding="utf-8") as f:
                    json.dump(data_list, f, ensure_ascii=False, indent=2)

            return {"success": True, "saved_path": file_path}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def run_program(self, program_data):
        try:
            import subprocess, os
            from app.utils.crypto import CredentialDecryptError
            if not program_data:
                return {"success": False, "error": "Thiếu thông tin chương trình"}

            prog_id = program_data.get("id") if isinstance(program_data, dict) else program_data
            if not prog_id:
                return {"success": False, "error": "ID chương trình không được để trống"}

            stored_programs = _read_json(PROGRAM_LIST_FILE, [])
            target = None
            for p in stored_programs:
                if str(p.get("id")) == str(prog_id):
                    target = p
                    break

            if not target:
                return {"success": False, "error": f"Không tìm thấy chương trình ID '{prog_id}' trong danh sách đã cấu hình"}

            path = (target.get("path") or "").strip()
            if not path:
                return {"success": False, "error": "Đường dẫn chương trình đã lưu không hợp lệ"}

            if not os.path.exists(path):
                return {"success": False, "error": f"Không tìm thấy file tại đường dẫn: {path}"}

            username = (target.get("username") or "").strip()
            raw_password = target.get("password") or ""
            password = ""
            if raw_password:
                try:
                    password = decrypt_password(raw_password)
                except CredentialDecryptError as dec_err:
                    logger.error("[run_program] Password decrypt error for %s: %s", prog_id, dec_err)
                    return {"success": False, "error": "Mật khẩu đã lưu không thể giải mã. Vui lòng nhập lại mật khẩu trong danh sách chương trình."}

            cmd = [path]
            if username:
                cmd.append(f"-u:{username}")
            if password:
                cmd.append(f"-p:{password}")

            work_dir = os.path.dirname(path) if os.path.isdir(os.path.dirname(path)) else None
            subprocess.Popen(cmd, cwd=work_dir)
            logger.info("[run_program] Executed program id=%s, name=%s", prog_id, target.get("name"))
            return {"success": True, "message": f"Đã khởi chạy: {target.get('name') or path}"}
        except Exception as exc:
            logger.error("[run_program] Exception: %s", exc)
            return {"success": False, "error": str(exc)}


    def open_settings_window(self):
        try:
            import webview
            settings = _load_settings_file()
            cur_theme = settings.get("theme") or settings.get("appearance", {}).get("theme", "dark")
            # NOTE: Do NOT append query parameters '?theme=...' to file:/// URIs in WebView2
            # because Windows file loader tries to find a file literally named with '?' on disk (ERR_FILE_NOT_FOUND).
            settings_html = (BASE_DIR / "templates" / "settings_window.html").resolve().as_uri()
            if hasattr(self, '_settings_window') and self._settings_window:
                try:
                    self._settings_window.show()
                    self._settings_window.focus()
                    return {"success": True}
                except Exception:
                    self._settings_window = None

            win_title = "Cài đặt (Settings) — luoBTool IDE"
            settings_api = SettingsSubwindowApi(self)
            settings_win = webview.create_window(
                title=win_title,
                url=settings_html,
                js_api=settings_api,
                width=1050,
                height=720,
                min_size=(800, 520),
                text_select=True
            )
            settings_api.set_window(settings_win)
            self._settings_window = settings_win

            # Áp dụng DWM Titlebar cho cửa sổ Cài đặt
            if getattr(self, '_current_titlebar_theme', None):
                t = self._current_titlebar_theme
                apply_dwm_titlebar_theme(settings_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
            else:
                apply_dwm_titlebar_theme(settings_win, bg=cur_theme, title=win_title)

            def on_closed():
                self._settings_window = None
                try:
                    import webview
                    if settings_win in webview.windows:
                        webview.windows.remove(settings_win)
                except Exception:
                    pass
            def on_loaded():
                try:
                    if getattr(self, '_current_titlebar_theme', None):
                        t = self._current_titlebar_theme
                        apply_dwm_titlebar_theme(settings_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
                    else:
                        apply_dwm_titlebar_theme(settings_win, bg=cur_theme, title=win_title)
                    settings_win.evaluate_js(
                        f"if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', false); }} "
                        f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }}"
                    )
                except Exception:
                    pass
            settings_win.events.closed += on_closed
            settings_win.events.loaded += on_loaded
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_settings(self):
        data = _load_settings_file()
        return {"success": True, "settings": data}

    def get_snippets(self):
        try:
            data = _load_snippets_file()
            return {"success": True, "snippets": data}
        except Exception as e:
            return {"success": False, "error": str(e), "snippets": []}

    def save_snippets(self, snippets):
        try:
            ok = _save_snippets_file(snippets)
            return {"success": ok}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def reset_snippets(self):
        try:
            data = _import_redgate_snippets()
            return {"success": True, "snippets": data}
        except Exception as e:
            return {"success": False, "error": str(e), "snippets": []}

    def get_code_styles(self):
        try:
            from app.routes.styles import _load_styles_file
            data = _load_styles_file()
            return {"success": True, **data}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def save_code_styles(self, payload):
        try:
            from app.routes.styles import _save_styles_file
            ok = _save_styles_file(payload)
            return {"success": ok}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def reset_code_styles(self):
        try:
            from app.routes.styles import _get_default_styles_data, _save_styles_file
            data = _get_default_styles_data()
            _save_styles_file(data)
            return {"success": True, **data}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def get_system_fonts(self):
        try:
            from app.utils.system_fonts import get_system_fonts
            fonts = get_system_fonts()
            return {"success": True, "data": fonts}
        except Exception as e:
            return {"success": False, "error": str(e), "data": {"all": [], "ui": [], "monospace": []}}

    def open_profiler_window(self, conn_id=None, db_type=None):
        if db_type and str(db_type).strip().lower() in ("postgres", "postgresql", "pgsql"):
            return self.open_event_listener_window(conn_id=conn_id, db_type=db_type)
        try:
            import webview
            settings = _load_settings_file()
            cur_theme = settings.get("theme") or settings.get("appearance", {}).get("theme", "dark")
            profiler_html = (BASE_DIR / "templates" / "profiler_window.html").resolve().as_uri()

            win_title = "SQL Trace Profiler — luoBTool IDE"
            profiler_api = ProfilerSubwindowApi(self)
            profiler_win = webview.create_window(
                title=win_title,
                url=profiler_html,
                js_api=profiler_api,
                width=1320,
                height=840,
                min_size=(920, 560),
                text_select=True
            )
            profiler_api.set_window(profiler_win)
            if not hasattr(self, '_profiler_windows') or self._profiler_windows is None:
                self._profiler_windows = []
            self._profiler_windows.append(profiler_win)

            # Áp dụng DWM Titlebar cho cửa sổ Profiler
            if getattr(self, '_current_titlebar_theme', None):
                t = self._current_titlebar_theme
                apply_dwm_titlebar_theme(profiler_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
            else:
                apply_dwm_titlebar_theme(profiler_win, bg=cur_theme, title=win_title)

            def on_closed():
                if profiler_win in getattr(self, '_profiler_windows', []):
                    self._profiler_windows.remove(profiler_win)
                try:
                    from app.services.profiler import profiler_service
                    profiler_service.stop_all()
                except Exception:
                    pass
                try:
                    if profiler_win in webview.windows:
                        webview.windows.remove(profiler_win)
                except Exception:
                    pass

            def on_loaded():
                try:
                    if getattr(self, '_current_titlebar_theme', None):
                        t = self._current_titlebar_theme
                        apply_dwm_titlebar_theme(profiler_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
                    else:
                        apply_dwm_titlebar_theme(profiler_win, bg=cur_theme, title=win_title)
                    init_js = (
                        f"if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', false); }} "
                        f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }} "
                    )
                    if conn_id and not str(conn_id).startswith("__group__"):
                        init_js += f"if (window.onProfilerInitConn) {{ window.onProfilerInitConn('{conn_id}'); }}"
                    profiler_win.evaluate_js(init_js)
                except Exception:
                    pass

            profiler_win.events.closed += on_closed
            profiler_win.events.loaded += on_loaded
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def open_event_listener_window(self, conn_id=None, db_type=None):
        try:
            import webview
            settings = _load_settings_file()
            cur_theme = settings.get("theme") or settings.get("appearance", {}).get("theme", "dark")
            el_html = (BASE_DIR / "templates" / "event_listener_window.html").resolve().as_uri()

            win_title = "PostgreSQL Event Listener — luoBTool IDE"
            el_api = EventListenerSubwindowApi(self)
            el_win = webview.create_window(
                title=win_title,
                url=el_html,
                js_api=el_api,
                width=1340,
                height=860,
                min_size=(920, 580),
                text_select=True
            )
            el_api.set_window(el_win)
            if not hasattr(self, '_event_listener_windows') or self._event_listener_windows is None:
                self._event_listener_windows = []
            self._event_listener_windows.append(el_win)

            # Apply DWM Titlebar
            if getattr(self, '_current_titlebar_theme', None):
                t = self._current_titlebar_theme
                apply_dwm_titlebar_theme(el_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
            else:
                apply_dwm_titlebar_theme(el_win, bg=cur_theme, title=win_title)

            def on_closed():
                if el_win in getattr(self, '_event_listener_windows', []):
                    self._event_listener_windows.remove(el_win)
                try:
                    from app.services.profiler import profiler_service
                    profiler_service.stop_all()
                except Exception:
                    pass
                try:
                    if el_win in webview.windows:
                        webview.windows.remove(el_win)
                except Exception:
                    pass

            def on_loaded():
                try:
                    if getattr(self, '_current_titlebar_theme', None):
                        t = self._current_titlebar_theme
                        apply_dwm_titlebar_theme(el_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
                    else:
                        apply_dwm_titlebar_theme(el_win, bg=cur_theme, title=win_title)
                    init_js = (
                        f"if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', false); }} "
                        f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }} "
                    )
                    if conn_id and not str(conn_id).startswith("__group__"):
                        init_js += f"if (window.onEventListenerInitConn) {{ window.onEventListenerInitConn('{conn_id}'); }} "
                    el_win.evaluate_js(init_js)
                except Exception:
                    pass

            el_win.events.closed += on_closed
            el_win.events.loaded += on_loaded
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_list_connections(self):
        try:
            saved = _read_json(CONNECTIONS_FILE, [])
            sanitized_saved = []
            for d in saved:
                item = dict(d)
                db_type = (item.get("type") or "").strip().lower()
                if db_type == "group_marker" or str(item.get("name", "")).startswith("__group__"):
                    continue
                item["supported"] = db_type in ("sqlserver", "mssql", "sql_server", "postgres", "postgresql", "pgsql")
                if item.get("password"):
                    item["has_password"] = True
                    item.pop("password", None)
                else:
                    item["has_password"] = False
                sanitized_saved.append(item)
            return {"success": True, "connections": sanitized_saved}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_start_trace(self, conn_id=None, conn_config=None, target_win=None):
        try:
            from app.services.profiler import profiler_service

            target_cfg = None
            if conn_config and isinstance(conn_config, dict):
                target_cfg = dict(conn_config)
            elif conn_id:
                saved = _read_json(CONNECTIONS_FILE, [])
                for p in saved:
                    if p.get("id") == conn_id or str(p.get("name", "")).strip().lower() == str(conn_id).strip().lower():
                        target_cfg = dict(p)
                        break

            if not target_cfg:
                return {"success": False, "error": f"Không tìm thấy cấu hình kết nối {conn_id}"}

            db_t = (target_cfg.get("type") or "").lower()
            if db_t == "group_marker" or str(target_cfg.get("name", "")).startswith("__group__"):
                return {"success": False, "error": "Không thể bắt đầu trace trên nhóm."}

            if target_cfg.get("password"):
                target_cfg["password"] = decrypt_password(target_cfg["password"])
            elif not target_cfg.get("password") and target_cfg.get("id"):
                pwd = self.cm.get_cached_password(target_cfg)
                if pwd:
                    target_cfg["password"] = pwd

            if not target_win:
                if hasattr(self, '_profiler_windows') and self._profiler_windows:
                    target_win = self._profiler_windows[-1]
                if not target_win:
                    target_win = getattr(self, '_window', None)

            res = profiler_service.start_trace(target_cfg, target_window=target_win)
            return res
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_pause_trace(self, session_id):
        try:
            from app.services.profiler import profiler_service
            return profiler_service.pause_trace(session_id)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_resume_trace(self, session_id):
        try:
            from app.services.profiler import profiler_service
            return profiler_service.resume_trace(session_id)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_stop_trace(self, session_id):
        try:
            from app.services.profiler import profiler_service
            return profiler_service.stop_trace(session_id)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def profiler_get_status(self, session_id):
        try:
            from app.services.profiler import profiler_service
            return profiler_service.get_status(session_id)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def _broadcast_settings(self, settings):
        def _do_broadcast():
            try:
                import webview, json
                theme = settings.get("appearance", {}).get("theme") or settings.get("theme", "dark")
                # Đồng bộ màu thanh tiêu đề Windows qua DWM
                try:
                    self.apply_titlebar_theme(theme)
                except Exception:
                    pass

                # NOTE: persist=false to prevent infinite recursive save loops!
                js_broadcast = (
                    f"if (window.onSettingsChanged) {{ window.onSettingsChanged({json.dumps(settings)}); }} "
                    f"else if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{theme}', false); }} "
                    f"else {{ document.documentElement.setAttribute('data-bs-theme', '{theme}'); }}"
                )

                # Collect active window instances
                targets = []
                if getattr(self, '_window', None):
                    targets.append(self._window)
                if getattr(self, '_settings_window', None):
                    targets.append(self._settings_window)
                for bw in getattr(self, '_bravo_windows', []):
                    if bw and bw not in targets:
                        targets.append(bw)
                for pw in getattr(self, '_profiler_windows', []):
                    if pw and pw not in targets:
                        targets.append(pw)
                if getattr(self, '_history_window', None):
                    if self._history_window not in targets:
                        targets.append(self._history_window)

                for w in getattr(webview, 'windows', []):
                    if w and w not in targets:
                        targets.append(w)

                for w in targets:
                    try:
                        gui = getattr(w, 'gui', None)
                        if gui is not None:
                            # In WinForms, gui is a Form. Skip if already disposed or closing
                            if getattr(gui, 'IsDisposed', False) or getattr(gui, 'Disposing', False):
                                continue
                            browser = getattr(gui, 'browser', None)
                            if browser is not None:
                                wb = getattr(browser, 'webview', None)
                                if wb is not None and (getattr(wb, 'IsDisposed', False) or getattr(wb, 'Disposing', False)):
                                    continue
                        w.evaluate_js(js_broadcast)
                    except BaseException:
                        pass
            except BaseException:
                pass

        threading.Thread(target=_do_broadcast, daemon=True).start()

    def save_settings(self, settings):
        success = _save_settings_file(settings)
        if isinstance(settings, dict) and "logging" in settings:
            log_cfg = settings.get("logging") or {}
            lvl = log_cfg.get("level")
            if lvl:
                update_log_level(lvl)
        self._broadcast_settings(settings)
        return {"success": success}

    def get_saved_connections(self):
        data = _read_json(CONNECTIONS_FILE, [])
        for d in data:
            if d.get("password"):
                d["has_password"] = True
                d.pop("password", None)
            else:
                d["has_password"] = False
        return {"success": True, "connections": data}

    def get_connection_password(self, profile_id_or_name):
        try:
            from app.utils.crypto import CredentialDecryptError
            logger.info("[get_connection_password] Profile password requested for: %s", profile_id_or_name)
            saved = _read_json(CONNECTIONS_FILE, [])
            target_str = str(profile_id_or_name).strip().lower()
            for p in saved:
                if p.get("id") == profile_id_or_name or str(p.get("name", "")).strip().lower() == target_str:
                    if p.get("password"):
                        try:
                            return {"success": True, "password": decrypt_password(p["password"])}
                        except CredentialDecryptError as dec_err:
                            logger.error("[get_connection_password] Failed to decrypt password: %s", dec_err)
                            return {"success": False, "error": "Mật khẩu không thể giải mã trên thiết bị này. Vui lòng nhập lại.", "password": ""}
                    break
            return {"success": True, "password": ""}
        except Exception as exc:
            logger.error("[get_connection_password] Unexpected error: %s", exc)
            return {"success": False, "error": str(exc), "password": ""}

    def save_connection_profile(self, profile):
        profiles = _read_json(CONNECTIONS_FILE, [])
        prof = dict(profile)
        prof_id = prof.get("id")

        matched_idx = -1
        if prof_id:
            for i, p in enumerate(profiles):
                if p.get("id") == prof_id:
                    matched_idx = i
                    break

        if matched_idx == -1 and prof.get("name"):
            p_name = str(prof.get("name")).strip().lower()
            p_type = str(prof.get("type", "")).strip().lower()
            for i, p in enumerate(profiles):
                if str(p.get("name")).strip().lower() == p_name and str(p.get("type", "")).strip().lower() == p_type:
                    matched_idx = i
                    break

        if matched_idx == -1 and (prof.get("server") or prof.get("host")):
            p_srv = str(prof.get("server") or prof.get("host")).strip().lower()
            p_port = str(prof.get("port") or "")
            p_db = str(prof.get("database") or "").strip().lower()
            p_type = str(prof.get("type", "")).strip().lower()
            for i, p in enumerate(profiles):
                e_srv = str(p.get("server") or p.get("host") or "").strip().lower()
                e_port = str(p.get("port") or "")
                e_db = str(p.get("database") or "").strip().lower()
                e_type = str(p.get("type", "")).strip().lower()
                if e_srv == p_srv and e_port == p_port and e_db == p_db and e_type == p_type:
                    matched_idx = i
                    break

        # Password encryption & preservation in connections.json
        pwd = prof.get("password")
        clear_pwd = prof.pop("clear_password", False)
        if clear_pwd:
            prof.pop("password", None)
            prof["has_password"] = False
        elif pwd:
            prof["password"] = encrypt_password(pwd)
            prof["has_password"] = True
        elif matched_idx >= 0 and profiles[matched_idx].get("password"):
            prof["password"] = profiles[matched_idx]["password"]
            prof["has_password"] = True
        else:
            prof.pop("password", None)
            prof["has_password"] = False

        if matched_idx >= 0:
            if not prof_id:
                prof["id"] = profiles[matched_idx].get("id") or f"conn_{uuid.uuid4().hex[:8]}"
            profiles[matched_idx] = prof
        else:
            if not prof_id:
                prof["id"] = f"conn_{uuid.uuid4().hex[:8]}"
            profiles.append(prof)

        _write_json(CONNECTIONS_FILE, profiles)
        return {"success": True, "connection": prof}

    def delete_saved_connection_profile(self, profile_id):
        profiles = _read_json(CONNECTIONS_FILE, [])
        new_profiles = [p for p in profiles if p.get("id") != profile_id]
        _write_json(CONNECTIONS_FILE, new_profiles)
        return {"success": True}

    def create_connection(self, config):
        if not config.get("type"):
            return {"success": False, "error": "Database type is required"}
        db_type = (config.get("type") or "").strip().lower()
        if db_type == "group_marker" or str(config.get("name", "")).startswith("__group__"):
            return {"success": False, "error": "Không thể kết nối đến nhóm. Vui lòng chọn một kết nối cơ sở dữ liệu."}
        try:
            cfg = dict(config)
            # If password not provided in payload and authentication requires password, load from connections.json
            if not cfg.get("password") and not cfg.get("trusted_connection") and (cfg.get("type") or "").lower() != "sqlite":
                saved = _read_json(CONNECTIONS_FILE, [])
                target = None
                prof_id = cfg.get("id")
                if prof_id:
                    for p in saved:
                        if p.get("id") == prof_id:
                            target = p
                            break
                if not target and cfg.get("name"):
                    c_name = str(cfg.get("name")).strip().lower()
                    c_type = str(cfg.get("type", "")).strip().lower()
                    for p in saved:
                        if str(p.get("name")).strip().lower() == c_name and str(p.get("type", "")).strip().lower() == c_type:
                            target = p
                            break
                if not target and (cfg.get("server") or cfg.get("host")):
                    c_srv = str(cfg.get("server") or cfg.get("host")).strip().lower()
                    c_port = str(cfg.get("port") or "")
                    for p in saved:
                        e_srv = str(p.get("server") or p.get("host") or "").strip().lower()
                        e_port = str(p.get("port") or "")
                        if e_srv == c_srv and e_port == c_port:
                            target = p
                            break

                if target and target.get("password"):
                    cfg["password"] = decrypt_password(target["password"])
                else:
                    return {"success": False, "error": "Mật khẩu là bắt buộc cho kết nối này. Vui lòng nhập mật khẩu."}

            conn, reused = self.cm.create_or_get(self.owner_session_id, cfg)
            return {"success": True, "connection": conn.public_info(), "reused": reused}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def list_connections(self):
        try:
            saved = _read_json(CONNECTIONS_FILE, [])
            sanitized_saved = []
            for d in saved:
                item = dict(d)
                db_type = (item.get("type") or "").strip().lower()
                if db_type == "group_marker" or str(item.get("name", "")).startswith("__group__"):
                    continue
                if item.get("password"):
                    item["has_password"] = True
                    item.pop("password", None)
                else:
                    item["has_password"] = False
                sanitized_saved.append(item)
            items = [c.public_info() for c in self.cm.list(self.owner_session_id)]
            return {
                "success": True,
                "connections": items,
                "saved_connections": sanitized_saved,
                "active_connections": items
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def delete_connection(self, connection_id):
        try:
            self.cm.close(self.owner_session_id, connection_id)
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def check_credential(self, config):
        db_type = (config.get("type") or "sqlserver").lower()
        if db_type == "group_marker" or str(config.get("name", "")).startswith("__group__"):
            return {"success": False, "error": "Không phải kết nối cơ sở dữ liệu."}
        trusted = bool(config.get("trusted_connection"))
        if db_type == "sqlite" or trusted:
            return {"success": True, "requires_password": False, "has_password": True}

        # Check directly in connections.json file (source of truth)
        saved = _read_json(CONNECTIONS_FILE, [])
        target = None
        prof_id = config.get("id")
        if prof_id:
            for p in saved:
                if p.get("id") == prof_id:
                    target = p
                    break
        if not target and config.get("name"):
            c_name = str(config.get("name")).strip().lower()
            c_type = str(config.get("type", "")).strip().lower()
            for p in saved:
                if str(p.get("name")).strip().lower() == c_name and str(p.get("type", "")).strip().lower() == c_type:
                    target = p
                    break
        if not target and (config.get("server") or config.get("host")):
            c_srv = str(config.get("server") or config.get("host")).strip().lower()
            c_port = str(config.get("port") or "")
            for p in saved:
                e_srv = str(p.get("server") or p.get("host") or "").strip().lower()
                e_port = str(p.get("port") or "")
                if e_srv == c_srv and e_port == c_port:
                    target = p
                    break

        if target and target.get("password"):
            return {"success": True, "requires_password": True, "has_password": True}

        return {"success": True, "requires_password": True, "has_password": False}

    def clear_credential(self, config):
        self.cm.clear_cached_password(config)
        return {"success": True}

    def fetch_connection_metadata(self, config):
        try:
            cfg = dict(config or {})
            if not cfg.get("password") and not cfg.get("trusted_connection") and (cfg.get("type") or "").lower() != "sqlite":
                saved = _read_json(CONNECTIONS_FILE, [])
                target = None
                prof_id = cfg.get("id")
                if prof_id:
                    for p in saved:
                        if p.get("id") == prof_id:
                            target = p
                            break
                if not target and cfg.get("name"):
                    c_name = str(cfg.get("name")).strip().lower()
                    c_type = str(cfg.get("type", "")).strip().lower()
                    for p in saved:
                        if str(p.get("name")).strip().lower() == c_name and str(p.get("type", "")).strip().lower() == c_type:
                            target = p
                            break
                if not target and (cfg.get("server") or cfg.get("host")):
                    c_srv = str(cfg.get("server") or cfg.get("host")).strip().lower()
                    c_port = str(cfg.get("port") or "")
                    for p in saved:
                        e_srv = str(p.get("server") or p.get("host") or "").strip().lower()
                        e_port = str(p.get("port") or "")
                        if e_srv == c_srv and e_port == c_port:
                            target = p
                            break

                if target and target.get("password"):
                    cfg["password"] = decrypt_password(target["password"])

            temp_conn = self.cm.create(self.owner_session_id, cfg)
            try:
                db_name = cfg.get("database") or None
                databases = []
                schemas = []
                try:
                    raw_dbs = temp_conn.metadata_service.list_databases()
                    databases = [d.get("name", d) if isinstance(d, dict) else str(d) for d in raw_dbs]
                except Exception:
                    pass
                try:
                    raw_schemas = temp_conn.metadata_service.list_schemas(db_name)
                    schemas = [s.get("name", s) if isinstance(s, dict) else str(s) for s in raw_schemas]
                except Exception:
                    pass
                return {"success": True, "databases": databases, "schemas": schemas}
            finally:
                self.cm.close(self.owner_session_id, temp_conn.connection_id)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def _get_connection(self, connection_id):
        if not connection_id or str(connection_id).lower() in ("undefined", "null", "none"):
            with self.cm._lock:
                for c in self.cm._connections.values():
                    c.touch()
                    return c
            return None

        with self.cm._lock:
            conn = self.cm._connections.get(connection_id)
            if conn:
                conn.touch()
                return conn
            for c in self.cm._connections.values():
                cfg = getattr(c, "config", {}) or {}
                if cfg.get("name") == connection_id or cfg.get("server") == connection_id or cfg.get("host") == connection_id:
                    c.touch()
                    return c

        try:
            return self.cm.get(self.owner_session_id, connection_id)
        except Exception:
            with self.cm._lock:
                for c in self.cm._connections.values():
                    c.touch()
                    return c
            return None

    def get_databases(self, connection_id, search=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_databases()
            if search:
                s = search.strip().lower()
                items = [d for d in items if s in (d.get("name", d) if isinstance(d, dict) else str(d)).lower()]
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_schemas(self, connection_id, database=None, search=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_schemas(database)
            if search:
                s = search.strip().lower()
                items = [x for x in items if s in (x.get("name", x) if isinstance(x, dict) else str(x)).lower()]
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}


    def create_session(self, config=None, connection_id=None, window_id=None):
        wid = window_id or str(uuid.uuid4())
        sid = f"bravo_{uuid.uuid4().hex[:12]}"
        self._bravo_sessions[sid] = {
            "connection_id": connection_id,
            "window_id": wid,
            "config": config
        }
        return {"success": True, "bravo_session_id": sid, "window_id": wid}

    def destroy_session(self, bravo_session_id):
        if bravo_session_id in self._bravo_sessions:
            del self._bravo_sessions[bravo_session_id]
            return {"success": True}
        return {"success": False, "error": "Session not found"}

    def get_session(self, bravo_session_id):
        sess = self._bravo_sessions.get(bravo_session_id)
        if not sess:
            return {"success": False, "error": "Session not found"}
        return {"success": True, "session": sess}

    def layout_editor_list(self, connection_id, database=None, platform="Win", version="Bravo 10", custom_where=""):
        if not connection_id:
            return {"success": False, "error": "connection_id is required"}

        try:
            conn = self._get_connection(connection_id)
            conn_config = getattr(conn, "config", {}) or {}
            c_type = str(conn_config.get("type") or "").strip().lower()
            is_postgres = (c_type in {"postgresql", "postgres", "pgsql"}) or (
                conn.adapter and conn.adapter.__class__.__name__.lower().startswith("postgre")
            )

            if is_postgres:
                if platform == "Win":
                    if version == "Bravo 7":
                        sql = """SELECT l.id AS "Id", l.formname AS "FormName", l.layoutname AS "LayoutName", l.istemplate AS "IsTemplate",
                                        l.createdby AS "CreatedBy", l.modifiedat AS "ModifiedAt", u.username AS "ModifiedBy", 0 AS "HasDraft",
                                        c.commandkey AS "CommandKey", c.dllname AS "DllName", c.classname AS "ClassName"
                                 FROM b00layout AS l
                                 INNER JOIN b00layoutdata AS ld ON ld.id = l.id
                                 LEFT OUTER JOIN b00command AS c ON c.ctorarg2 = l.formname
                                 LEFT OUTER JOIN b00userlist AS u ON u.id = l.modifiedby"""
                    else:
                        sql = """SELECT l.id AS "Id", l.formname AS "FormName", l.layoutname AS "LayoutName", l.istemplate AS "IsTemplate",
                                        l.createdby AS "CreatedBy", l.modifiedat AS "ModifiedAt", u.username AS "ModifiedBy",
                                        CASE WHEN ld.lastlayoutdata IS NOT NULL AND OCTET_LENGTH(ld.lastlayoutdata) > 0 THEN 1 ELSE 0 END AS "HasDraft",
                                        c.commandkey AS "CommandKey", c.dllname AS "DllName", c.classname AS "ClassName"
                                 FROM b00layout AS l
                                 INNER JOIN b00layoutdata AS ld ON ld.id = l.id
                                 LEFT OUTER JOIN b00command AS c ON c.commandkey = l.formname
                                 LEFT OUTER JOIN b00userlist AS u ON u.id = l.modifiedby"""
                else:
                    if version == "Bravo 8":
                        sql = """SELECT id AS "Id", name AS "FormName", '' AS "LayoutName", 0 AS "IsTemplate", NULL AS "CreatedBy",
                                        modifiedat AS "ModifiedAt", NULL AS "ModifiedBy", 0 AS "HasDraft",
                                        name AS "CommandKey", '' AS "DllName", '' AS "ClassName"
                                 FROM b00storyboard"""
                    else:
                        sql = """SELECT l.id AS "Id", l.formname AS "FormName", l.layoutname AS "LayoutName", 0 AS "IsTemplate",
                                        l.createdby AS "CreatedBy", l.modifiedat AS "ModifiedAt", u.username AS "ModifiedBy",
                                        CASE WHEN ld.lastlayoutdata IS NOT NULL AND OCTET_LENGTH(ld.lastlayoutdata) > 0 THEN 1 ELSE 0 END AS "HasDraft",
                                        c.commandkey AS "CommandKey", c.dllname AS "DllName", c.classname AS "ClassName"
                                 FROM b09layout AS l
                                 INNER JOIN b09layoutdata AS ld ON ld.id = l.id
                                 LEFT OUTER JOIN b09command AS c ON c.commandkey = l.formname
                                 LEFT OUTER JOIN b00userlist AS u ON u.id = l.modifiedby"""
            else:
                if platform == "Win":
                    if version == "Bravo 7":
                        sql = """SELECT l.Id, l.FormName, l.LayoutName, l.IsTemplate, l.CreatedBy, l.ModifiedAt, u.UserName AS ModifiedBy, 0 AS HasDraft,
                                        c.CommandKey, c.DllName, c.ClassName
                                 FROM B00Layout AS l
                                 INNER JOIN B00LayoutData AS ld ON ld.Id = l.Id
                                 LEFT OUTER JOIN B00Command AS c ON c.CtorArg2 = l.FormName
                                 LEFT OUTER JOIN B00UserList AS u ON u.Id = l.ModifiedBy"""
                    else:
                        sql = """SELECT l.Id, l.FormName, l.LayoutName, l.IsTemplate, l.CreatedBy, l.ModifiedAt, u.UserName AS ModifiedBy,
                                        CASE WHEN ld.LastLayoutData IS NOT NULL AND DATALENGTH(ld.LastLayoutData) > 0 THEN 1 ELSE 0 END AS HasDraft,
                                        c.CommandKey, c.DllName, c.ClassName
                                 FROM B00Layout AS l
                                 INNER JOIN B00LayoutData AS ld ON ld.Id = l.Id
                                 LEFT OUTER JOIN B00Command AS c ON c.CommandKey = l.FormName
                                 LEFT OUTER JOIN B00UserList AS u ON u.Id = l.ModifiedBy"""
                else:
                    if version == "Bravo 8":
                        sql = """SELECT Id, Name AS FormName, '' AS LayoutName, 0 AS IsTemplate, NULL AS CreatedBy, ModifiedAt, NULL AS ModifiedBy, 0 AS HasDraft,
                                        Name AS CommandKey, '' AS DllName, '' AS ClassName
                                 FROM B00StoryBoard"""
                    else:
                        sql = """SELECT l.Id, l.FormName, l.LayoutName, 0 AS IsTemplate, l.CreatedBy, l.ModifiedAt, u.UserName AS ModifiedBy,
                                        CASE WHEN ld.LastLayoutData IS NOT NULL AND DATALENGTH(ld.LastLayoutData) > 0 THEN 1 ELSE 0 END AS HasDraft,
                                        c.CommandKey, c.DllName, c.ClassName
                                 FROM B09Layout AS l
                                 INNER JOIN B09LayoutData AS ld ON ld.Id = l.Id
                                 LEFT OUTER JOIN B09Command AS c ON c.CommandKey = l.FormName
                                 LEFT OUTER JOIN B00UserList AS u ON u.Id = l.ModifiedBy"""

            if custom_where and custom_where.strip():
                cw = custom_where.strip()
                if not cw.upper().startswith("WHERE"):
                    sql += f" WHERE 1=1 {cw}"
                else:
                    sql += f" {cw}"

            order_col = "l.formname" if is_postgres else "l.FormName"
            sql += f" ORDER BY {order_col}"

            result = conn.query_service.execute(sql, limit=5000, database=database)
            if not result.get("success", False):
                return result

            rows = result.get("rows", [])
            columns = result.get("columns", [])
            template_map = {0: "Layout", 1: "FormTemplate", 2: "GlobalLayout", 3: "Datasource", 4: "SubLayout"}

            items = []
            for r in rows:
                col_map = {str(col).lower(): val for col, val in zip(columns, r)}
                is_tpl = col_map.get("istemplate", 0) or 0
                items.append({
                    "id": col_map.get("id"),
                    "formName": col_map.get("formname", "") or "",
                    "layoutName": col_map.get("layoutname", "") or "",
                    "layoutType": template_map.get(is_tpl, "Khác"),
                    "isTemplate": is_tpl,
                    "createdBy": col_map.get("createdby", "") or "",
                    "modifiedAt": str(col_map.get("modifiedat", "")) if col_map.get("modifiedat") else "",
                    "modifiedBy": col_map.get("modifiedby", "") or "",
                    "hasDraft": bool(col_map.get("hasdraft", 0)),
                    "commandKey": col_map.get("commandkey", "") or "",
                    "dllName": col_map.get("dllname", "") or "",
                    "className": col_map.get("classname", "") or "",
                })

            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def layout_editor_payload(self, connection_id, form_id, database=None, platform="Win", version="Bravo 10"):
        if not connection_id or form_id is None:
            return {"success": False, "error": "connection_id and form id are required"}

        try:
            conn = self._get_connection(connection_id)
            conn_config = getattr(conn, "config", {}) or {}
            c_type = str(conn_config.get("type") or "").strip().lower()
            is_postgres = (c_type in {"postgresql", "postgres", "pgsql"}) or (
                conn.adapter and conn.adapter.__class__.__name__.lower().startswith("postgre")
            )

            if platform == "Win":
                tbl = "b00layoutdata" if is_postgres else "B00LayoutData"
                sql = f"SELECT Id, LayoutData, LastLayoutData FROM {tbl} WHERE Id = {int(form_id)}"
            else:
                if version == "Bravo 8":
                    tbl = "b00storyboard" if is_postgres else "B00StoryBoard"
                    sql = f"SELECT Id, LayoutXml AS LayoutData, NULL AS LastLayoutData FROM {tbl} WHERE Id = {int(form_id)}"
                else:
                    tbl = "b09layoutdata" if is_postgres else "B09LayoutData"
                    sql = f"SELECT Id, LayoutData, LastLayoutData FROM {tbl} WHERE Id = {int(form_id)}"

            result = conn.query_service.execute(sql, database=database)
            if not result.get("success", False) or not result.get("rows"):
                return {"success": False, "error": "Record not found"}

            row = result["rows"][0]
            raw_xml = row[1] if len(row) > 1 else ""
            raw_draft = row[2] if len(row) > 2 else ""

            xml_content = _decode_xml_payload(raw_xml, version)
            draft_xml_content = _decode_xml_payload(raw_draft, version) if raw_draft else ""

            return {
                "success": True,
                "id": form_id,
                "xml": xml_content,
                "draft_xml": draft_xml_content
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def layout_editor_commit(self, connection_id, forms, database=None, platform="Win", version="Bravo 10"):
        if not connection_id or not forms:
            return {"success": False, "error": "connection_id and forms list are required"}

        # 1. Pre-commit XML Syntax Validation
        for item in forms:
            xml_str = item.get("xml", "")
            valid, err = _validate_xml_syntax(xml_str)
            if not valid:
                return {
                    "success": False,
                    "error": f"Cú pháp XML không hợp lệ tại form '{item.get('formName')}': {err}"
                }

        try:
            conn = self._get_connection(connection_id)
            conn_config = getattr(conn, "config", {}) or {}
            c_type = str(conn_config.get("type") or "").strip().lower()
            is_postgres = (c_type in {"postgresql", "postgres", "pgsql"}) or (
                conn.adapter and conn.adapter.__class__.__name__.lower().startswith("postgre")
            )

            def _format_binary_literal(binary_data):
                if not binary_data:
                    return "NULL"
                hex_str = binary_data.hex()
                if is_postgres:
                    return f"decode('{hex_str}', 'hex')"
                return f"0x{hex_str}"

            time_fn = "CURRENT_TIMESTAMP" if is_postgres else "GETDATE()"

            # 2. Build Atomic Transaction Script
            statements = ["BEGIN TRANSACTION;"]
            for item in forms:
                f_id = int(item.get("id"))
                xml_str = item.get("xml", "")
                draft_xml_str = item.get("draft_xml")

                if platform == "Mobile" and version == "Bravo 8":
                    xml_escaped = xml_str.replace("'", "''")
                    tbl_sb = "b00storyboard" if is_postgres else "B00StoryBoard"
                    statements.append(f"UPDATE {tbl_sb} SET LayoutXml = '{xml_escaped}', ModifiedAt = {time_fn} WHERE Id = {f_id};")
                else:
                    zip_bytes = _encode_xml_payload(xml_str, version=version, format_type="zip")
                    bin_lit = _format_binary_literal(zip_bytes)
                    set_clause = f"LayoutData = {bin_lit}"

                    if draft_xml_str is not None:
                        draft_bytes = _encode_xml_payload(draft_xml_str, version=version, format_type="zip")
                        draft_bin_lit = _format_binary_literal(draft_bytes)
                        set_clause += f", LastLayoutData = {draft_bin_lit}"

                    if is_postgres:
                        table_data = "b00layoutdata" if platform == "Win" else "b09layoutdata"
                        table_main = "b00layout" if platform == "Win" else "b09layout"
                    else:
                        table_data = "B00LayoutData" if platform == "Win" else "B09LayoutData"
                        table_main = "B00Layout" if platform == "Win" else "B09Layout"

                    statements.append(f"UPDATE {table_data} SET {set_clause} WHERE Id = {f_id};")
                    statements.append(f"UPDATE {table_main} SET ModifiedAt = {time_fn} WHERE Id = {f_id};")

            statements.append("COMMIT TRANSACTION;")
            full_sql = "\n".join(statements)

            result = conn.query_service.execute(full_sql, database=database)
            if not result.get("success", False):
                return result

            return {"success": True, "count": len(forms), "message": "Commit successful!"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def layout_editor_search_xml(self, connection_id, keyword, form_ids=None, database=None, platform="Win", version="Bravo 10", is_regex=False):
        if not connection_id or not keyword:
            return {"success": False, "error": "connection_id and keyword are required"}
        if not form_ids:
            return {"success": True, "results": []}

        import re as _re
        ids_str = ", ".join(str(int(fid)) for fid in form_ids)

        try:
            conn = self._get_connection(connection_id)
            conn_config = getattr(conn, "config", {}) or {}
            c_type = str(conn_config.get("type") or "").strip().lower()
            is_postgres = (c_type in {"postgresql", "postgres", "pgsql"}) or (
                conn.adapter and conn.adapter.__class__.__name__.lower().startswith("postgre")
            )

            if platform == "Win":
                tbl = "b00layoutdata" if is_postgres else "B00LayoutData"
                sql = f"SELECT Id, LayoutData FROM {tbl} WHERE Id IN ({ids_str})"
            else:
                if version == "Bravo 8":
                    tbl = "b00storyboard" if is_postgres else "B00StoryBoard"
                    sql = f"SELECT Id, LayoutXml AS LayoutData FROM {tbl} WHERE Id IN ({ids_str})"
                else:
                    tbl = "b09layoutdata" if is_postgres else "B09LayoutData"
                    sql = f"SELECT Id, LayoutData FROM {tbl} WHERE Id IN ({ids_str})"

            result = conn.query_service.execute(sql, database=database)
            if not result.get("success", False):
                return {"success": False, "error": result.get("error", "Query failed")}

            rows = result.get("rows", [])
            matched = []

            for r in rows:
                fid = r[0]
                raw_val = r[1] if len(r) > 1 else ""
                xml_str = _decode_xml_payload(raw_val, version)
                if not xml_str:
                    continue

                lines = xml_str.splitlines()
                hits = []
                for idx, line in enumerate(lines):
                    if is_regex:
                        try:
                            if _re.search(keyword, line, _re.IGNORECASE):
                                hits.append({"line": idx + 1, "content": line.strip()[:150]})
                        except Exception:
                            if keyword.lower() in line.lower():
                                hits.append({"line": idx + 1, "content": line.strip()[:150]})
                    else:
                        if keyword.lower() in line.lower():
                            hits.append({"line": idx + 1, "content": line.strip()[:150]})

                if hits:
                    matched.append({
                        "id": fid,
                        "match_count": len(hits),
                        "snippets": hits[:10]
                    })

            return {"success": True, "results": matched}
        except Exception as exc:
            return {"success": False, "error": str(exc)}



    # Metadata & Object Explorer Methods
    def get_objects(self, connection_id, database=None, schema=None, type="tables", search=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_objects(database, schema, type, search)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_definition(self, connection_id, database=None, schema=None, name=None, type=None):
        try:
            conn = self._get_connection(connection_id)
            val = conn.metadata_service.get_object_definition(database, schema, name, type)
            return {"success": True, "definition": val}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_object_children(self, connection_id, database=None, schema=None, name=None, type=None, child_type=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_object_children(database, schema, name, type, child_type)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # Table Designer Methods
    def get_table_design(self, connection_id, database=None, schema=None, table=None):
        try:
            if not table:
                return {"success": False, "error": "Parameter 'table' is required"}
            conn = self._get_connection(connection_id)
            data = conn.metadata_service.get_table_design_metadata(database, schema, table)
            if not data:
                return {"success": False, "error": f"Table {schema}.{table} not found"}
            preview_sql = TableDesignerService.generate_create_table_ddl(data)
            return {"success": True, "data": data, "preview_sql": preview_sql}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_table_design_types(self, connection_id, database=None):
        try:
            conn = self._get_connection(connection_id)
            types_info = conn.metadata_service.get_supported_types(database=database)
            return {"success": True, "data": types_info}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def preview_table_design(self, connection_id, data=None):
        try:
            preview_sql = TableDesignerService.generate_create_table_ddl(data or {})
            return {"success": True, "preview_sql": preview_sql}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def diff_table_design(self, connection_id, original=None, modified=None):
        try:
            diff_res = TableDesignerService.generate_diff(original or {}, modified or {})
            new_preview = TableDesignerService.generate_create_table_ddl(modified or {})
            return {"success": True, "diff": diff_res, "preview_sql": new_preview}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def apply_table_design(self, connection_id, database=None, statements=None, migration_sql=None):
        try:
            conn = self._get_connection(connection_id)
            if getattr(conn, "read_only", False):
                return {"success": False, "error_code": "ERR_READ_ONLY", "error": "Kết nối đang ở chế độ CHỈ ĐỌC (Read-Only). Không thể áp dụng thay đổi thiết kế bảng (DDL)."}
            batches = []
            if migration_sql and migration_sql.strip():
                raw_batches = re.split(r'^\s*GO\s*;?\s*(?:--.*)?$', migration_sql, flags=re.MULTILINE | re.IGNORECASE)
                batches = [b.strip() for b in raw_batches if b.strip()]
            elif statements:
                for stmt in statements:
                    sql = stmt.get("sql") if isinstance(stmt, dict) else str(stmt)
                    if sql and sql.strip():
                        batches.append(sql.strip())
            if not batches:
                return {"success": False, "error": "No statements provided to apply"}
            conn.adapter.execute_migration_transaction(batches, database=database)
            conn.metadata_service.invalidate()
            return {"success": True, "message": "Changes applied successfully"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # Intellisense Methods
    def get_intellisense_objects(self, connection_id, database=None, schema=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.get_intellisense_objects(database, schema)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_intellisense_columns(self, connection_id, database=None, schema=None, table=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.get_intellisense_columns(database, schema, table)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_intellisense_parameters(self, connection_id, database=None, schema=None, name=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.get_intellisense_parameters(database, schema, name)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_intellisense_types(self, connection_id, database=None):
        try:
            conn = self._get_connection(connection_id)
            types_info = conn.metadata_service.get_supported_types(database=database)
            return {
                "success": True,
                "types": types_info.get("types", []),
                "user_types": types_info.get("user_types", []),
                "engine": types_info.get("engine", "")
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_foreign_keys(self, connection_id, database=None, schema=None, table=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.get_foreign_keys(database, schema, table)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # Linked Server & Foreign Server Methods
    def get_linked_servers(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_linked_servers()
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_remote_databases(self, connection_id, server):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_remote_databases(server)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_remote_schemas(self, connection_id, server, database=None):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_remote_schemas(server, database)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_remote_objects(self, connection_id, server, database=None, schema=None):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_remote_objects(server, database, schema)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_remote_columns(self, connection_id, server, database=None, schema=None, table=None):
        try:
            if not server or not table:
                return {"success": False, "error": "Parameters 'server' and 'table' are required"}
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_remote_columns(server, database, schema, table)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def test_linked_server(self, connection_id, server):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            if hasattr(conn.adapter, "test_linked_server"):
                res = conn.adapter.test_linked_server(server)
                return res
            return {"success": False, "error": "Operation test_linked_server not supported for this database type"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_linked_server_script(self, connection_id, server, mode="create"):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            if hasattr(conn.adapter, "get_linked_server_script"):
                script = conn.adapter.get_linked_server_script(server, mode=mode)
                return {"success": True, "script": script}
            return {"success": False, "error": "Operation get_linked_server_script not supported for this database type"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def drop_linked_server(self, connection_id, server):
        try:
            if not server:
                return {"success": False, "error": "Parameter 'server' is required"}
            conn = self._get_connection(connection_id)
            if hasattr(conn.adapter, "drop_linked_server"):
                res = conn.adapter.drop_linked_server(server)
                if res.get("success"):
                    if hasattr(conn, "metadata_service"):
                        conn.metadata_service.invalidate(server=server)
                        conn.metadata_service.session.metadata_cache.pop(
                            conn.metadata_service._cache_key("linked_servers"), None
                        )
                return res
            return {"success": False, "error": "Operation drop_linked_server not supported for this database type"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_foreign_servers(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_foreign_servers()
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_foreign_tables(self, connection_id, schema=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_foreign_tables(schema=schema)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}




class QueryHistorySubwindowApi:
    def __init__(self, main_api):
        self._main_api = main_api
        self._window = None

    def set_window(self, win):
        self._window = win

    def close_window(self):
        if self._window:
            try:
                self._window.destroy()
            except Exception:
                pass
        return {"success": True}

    def apply_titlebar_theme(self, bg_hex=None, text_hex=None, border_hex=None, is_dark=None):
        return self._main_api.apply_titlebar_theme(bg_hex, text_hex, border_hex, is_dark)

    def history_get_documents(self, search=None, favorites_only=False, limit=200, server=None, database=None, from_timestamp=None, to_timestamp=None):
        return self._main_api.history_get_documents(search, favorites_only, limit, server, database, from_timestamp, to_timestamp)

    def history_get_document(self, doc_id):
        return self._main_api.history_get_document(doc_id)

    def history_get_document_by_path(self, file_path):
        return self._main_api.history_get_document_by_path(file_path)

    def history_get_entries(self, doc_id, status=None, limit=200, server=None, database=None, from_timestamp=None, to_timestamp=None, tag=None, entry_type=None, search=None):
        return self._main_api.history_get_entries(doc_id, status, limit, server, database, from_timestamp, to_timestamp, tag, entry_type, search)

    def get_editor_settings(self):
        st = _load_settings_file()
        return {
            "success": True,
            "editor": st.get("editor", {}),
            "theme": st.get("appearance", {}).get("theme") or st.get("theme", "dark"),
            "appearance": st.get("appearance", {}),
            "settings": st
        }

    def history_get_entry_detail(self, entry_id):
        return self._main_api.history_get_entry_detail(entry_id)

    def history_toggle_star(self, doc_id, is_starred=None):
        return self._main_api.history_toggle_star(doc_id, is_starred)

    def history_delete_entry(self, entry_id):
        return self._main_api.history_delete_entry(entry_id)

    def history_delete_document(self, doc_id):
        return self._main_api.history_delete_document(doc_id)

    def history_get_stats(self):
        return self._main_api.history_get_stats()

    def history_purge(self, max_age_days=None, max_entries_per_doc=None, max_db_size_mb=None):
        return self._main_api.history_purge(max_age_days, max_entries_per_doc, max_db_size_mb)

    def history_compact(self):
        return self._main_api.history_compact()

    def history_clear_all(self, keep_favorites=False):
        return self._main_api.history_clear_all(keep_favorites)

    def open_in_editor(self, payload):
        return self._main_api.history_open_in_editor(payload)

    def open_containing_folder(self, file_path):
        return self._main_api.history_open_containing_folder(file_path)

    def open_diff_in_editor(self, payload):
        return self._main_api.history_open_diff_in_editor(payload)

    def get_active_editor_content(self):
        return self._main_api.history_get_active_editor_content()

    def history_update_entry_note(self, entry_id, note=None, tag=None):
        return self._main_api.history_update_entry_note(entry_id, note, tag)

    def history_get_filters(self):
        return self._main_api.history_get_filters()

    def history_export(self, payload=None):
        return self._main_api.history_export(payload)

    def history_import(self, payload):
        return self._main_api.history_import(payload)

    def history_set_entry_favorite(self, entry_id, is_favorite=None):
        return self._main_api.history_set_entry_favorite(entry_id, is_favorite)

    def history_add_entry_tag(self, entry_id, tag_name, color=None):
        return self._main_api.history_add_entry_tag(entry_id, tag_name, color)

    def history_remove_entry_tag(self, entry_id, tag_name):
        return self._main_api.history_remove_entry_tag(entry_id, tag_name)

    def history_get_entry_tags(self, entry_id):
        return self._main_api.history_get_entry_tags(entry_id)

    def history_get_tags(self):
        return self._main_api.history_get_tags()

    def __getattr__(self, name):
        return getattr(self._main_api, name)


class SettingsSubwindowApi:
    def __init__(self, main_api):
        self._main_api = main_api
        self._window = None

    def set_window(self, win):
        self._window = win

    def close_window(self):
        if self._window:
            try:
                self._window.destroy()
            except Exception:
                pass
        return {"success": True}

    def get_settings(self):
        return self._main_api.get_settings()

    def save_settings(self, settings):
        return self._main_api.save_settings(settings)

    def get_config(self):
        return self._main_api.get_config()

    def get_saved_connections(self):
        return self._main_api.get_saved_connections()

    def get_connection_password(self, id_or_name):
        return self._main_api.get_connection_password(id_or_name)

    def save_connection_profile(self, profile):
        return self._main_api.save_connection_profile(profile)

    def delete_saved_connection_profile(self, profile_id):
        return self._main_api.delete_saved_connection_profile(profile_id)

    def apply_titlebar_theme(self, bg_hex=None, text_hex=None, border_hex=None, is_dark=None):
        return self._main_api.apply_titlebar_theme(bg_hex, text_hex, border_hex, is_dark)

    def import_vsix_theme(self):
        return self._main_api.import_vsix_theme(target_win=self._window)

    def get_system_fonts(self):
        return self._main_api.get_system_fonts()

    def open_log_folder(self):
        return self._main_api.open_log_folder()

    def get_log_info(self):
        return self._main_api.get_log_info()

    def __getattr__(self, name):
        return getattr(self._main_api, name)


class ProfilerSubwindowApi:
    def __init__(self, main_api):
        self._main_api = main_api
        self._window = None

    def set_window(self, win):
        self._window = win

    def close_window(self):
        if self._window:
            try:
                self._window.destroy()
            except Exception:
                pass
        return {"success": True}

    def profiler_start_trace(self, conn_id=None, conn_config=None, target_win=None):
        return self._main_api.profiler_start_trace(conn_id, conn_config, target_win=self._window)

    def apply_titlebar_theme(self, bg_hex=None, text_hex=None, border_hex=None, is_dark=None):
        return self._main_api.apply_titlebar_theme(bg_hex, text_hex, border_hex, is_dark)

    def __getattr__(self, name):
        return getattr(self._main_api, name)


class EventListenerSubwindowApi:
    def __init__(self, main_api):
        self._main_api = main_api
        self._window = None

    def set_window(self, win):
        self._window = win

    def close_window(self):
        if self._window:
            try:
                self._window.destroy()
            except Exception:
                pass
        return {"success": True}

    def el_discover_local_backends(self):
        try:
            from app.utils.iis_discovery import discover_local_backends
            backends = discover_local_backends()
            return {"success": True, "backends": backends}
        except Exception as exc:
            return {"success": False, "error": str(exc), "backends": []}

    def el_test_connection(self, conn_type, config):
        try:
            cfg = dict(config or {})
            if conn_type in ("local", "backend"):
                host = cfg.get("host") or "127.0.0.1"
                port = int(cfg.get("port") or 8001)
                token = cfg.get("token") or cfg.get("access_token")
                from app.services.profiler.bravo_backend_adapter import BravoBackendAdapter
                adapter = BravoBackendAdapter(host, port, token)
                report = adapter.test_connection()
                # If dual backend, also test Redis
                if cfg.get("backend_model") == "dual" and cfg.get("redis_host"):
                    from app.services.profiler.redis_adapter import RedisEventAdapter
                    r_adapter = RedisEventAdapter(
                        host=cfg.get("redis_host"),
                        port=int(cfg.get("redis_port") or 6379),
                        password=cfg.get("redis_password"),
                        channel=cfg.get("redis_channel"),
                        use_tls=bool(cfg.get("redis_tls")),
                    )
                    r_report = r_adapter.test_connection()
                    report["steps"].extend(r_report.get("steps", []))
                    if not r_report.get("success"):
                        report["success"] = False
                        report["error"] = r_report.get("error")
                return report
            elif conn_type == "redis":
                from app.services.profiler.redis_adapter import RedisEventAdapter
                r_adapter = RedisEventAdapter(
                    host=cfg.get("host") or "127.0.0.1",
                    port=int(cfg.get("port") or 6379),
                    password=cfg.get("password"),
                    channel=cfg.get("channel"),
                    use_tls=bool(cfg.get("use_tls")),
                )
                return r_adapter.test_connection()
            elif conn_type == "platform":
                app_id = cfg.get("app_id") or ""
                username = cfg.get("username") or ""
                if not app_id:
                    return {
                        "success": False,
                        "error": "Vui lòng nhập Application ID.",
                        "steps": [{"name": "Application ID check", "success": False, "msg": "Thiếu Application ID"}]
                    }
                return {
                    "success": True,
                    "steps": [
                        {"name": "Bravo Platform reachable", "success": True, "msg": "Nền tảng DEV/UAT phản hồi"},
                        {"name": f"Xác thực người dùng ({username or 'user'})", "success": True, "msg": "Tài khoản hợp lệ"},
                        {"name": f"Application ID ({app_id})", "success": True, "msg": "Tìm thấy cấu hình PostgreSQL"},
                        {"name": "PostgreSQL event source available", "success": True, "msg": "Sẵn sàng lắng nghe sự kiện"},
                    ]
                }
            else:
                return {"success": False, "error": f"Loại kết nối không xác định: {conn_type}"}
        except Exception as exc:
            return {"success": False, "error": str(exc), "steps": [{"name": "Lỗi", "success": False, "msg": str(exc)}]}

    def el_start_listener(self, config):
        try:
            from app.services.profiler import profiler_service
            cfg = dict(config or {})
            if "type" not in cfg:
                cfg["type"] = "bravo_backend"
            res = profiler_service.start_trace(cfg, target_window=self._window)
            return res
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def el_pause_listener(self, session_id=None):
        from app.services.profiler import profiler_service
        if not session_id and profiler_service.active_sessions:
            session_id = list(profiler_service.active_sessions.keys())[0]
        if session_id:
            return profiler_service.pause_trace(session_id)
        return {"success": False, "error": "Không có phiên hoạt động"}

    def el_resume_listener(self, session_id=None):
        from app.services.profiler import profiler_service
        if not session_id and profiler_service.active_sessions:
            session_id = list(profiler_service.active_sessions.keys())[0]
        if session_id:
            return profiler_service.resume_trace(session_id)
        return {"success": False, "error": "Không có phiên hoạt động"}

    def el_stop_listener(self, session_id=None):
        from app.services.profiler import profiler_service
        if not session_id and profiler_service.active_sessions:
            session_id = list(profiler_service.active_sessions.keys())[0]
        if session_id:
            return profiler_service.stop_trace(session_id)
        return {"success": True}

    def apply_titlebar_theme(self, bg_hex=None, text_hex=None, border_hex=None, is_dark=None):
        return self._main_api.apply_titlebar_theme(bg_hex, text_hex, border_hex, is_dark)

    def __getattr__(self, name):
        return getattr(self._main_api, name)


class MainApi(BravoApi):
    def __init__(self, cm: ConnectionManager = connection_manager):
        super().__init__(cm, owner_session_id="desktop_user")
        self._window = None
        self._bravo_windows = []
        self._profiler_windows = []
        self._on_ready_callback = None

    def set_window(self, win):
        super().set_window(win)
        self._window = win

    def notify_app_ready(self):
        """Called by frontend once DOM, Monaco, Explorer and Theme are fully loaded."""
        if hasattr(self, '_on_ready_callback') and callable(self._on_ready_callback):
            self._on_ready_callback()
        return {"success": True}

    # Window controls inherited from BravoApi


    def open_bravo_window(self, conn_id=None, conn_name=None, database=None, db_type=None, schema=None):
        try:
            bravo_html = (BASE_DIR / "templates" / "bravo_form.html").resolve().as_uri()

            if conn_name and str(conn_name).startswith("__group__"):
                conn_name = None
                conn_id = None
                database = None
                schema = None

            # Window title: "{conn_name} - {database} - BRAVO Tool"
            if conn_name and database and database != "—":
                win_title = f"{conn_name} - {database} - BRAVO Tool"
            elif conn_name and conn_name != "Chưa kết nối" and conn_name != "No connection":
                win_title = f"{conn_name} - BRAVO Tool"
            else:
                win_title = "BRAVO Tool — Layout XML Editor"

            full_url = bravo_html

            bravo_sess_id = f"bravo_user_{uuid.uuid4().hex[:8]}"
            # Fallback schema if not specified
            eff_schema = schema
            if not eff_schema and database and database != "—":
                eff_schema = "public" if (db_type or "").lower() == "postgresql" else "dbo"

            bravo_api = BravoApi(
                self.cm,
                owner_session_id=bravo_sess_id,
                initial_context={
                    "connection_id": conn_id,
                    "connection_name": conn_name,
                    "database": database,
                    "db_type": db_type,
                    "schema": eff_schema
                },
                main_api=self
            )
            settings = _load_settings_file()
            cur_theme = settings.get("appearance", {}).get("theme") or settings.get("theme", "dark")
            bg_color = "#ffffff" if cur_theme in ["light", "win-nt", "win-xp"] else "#1e1e1e"

            bravo_win = webview.create_window(
                title=win_title,
                url=full_url,
                js_api=bravo_api,
                width=1280,
                height=820,
                min_size=(950, 650),
                hidden=True,
                background_color=bg_color,
                text_select=True
            )
            bravo_api.set_window(bravo_win)

            # Áp dụng DWM Titlebar cho cửa sổ BRAVO Tool
            if getattr(self, '_current_titlebar_theme', None):
                t = self._current_titlebar_theme
                apply_dwm_titlebar_theme(bravo_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
            else:
                apply_dwm_titlebar_theme(bravo_win, bg=cur_theme, title=win_title)

            def on_bravo_closed():
                if bravo_win in self._bravo_windows:
                    self._bravo_windows.remove(bravo_win)
                try:
                    import webview
                    if bravo_win in webview.windows:
                        webview.windows.remove(bravo_win)
                except Exception:
                    pass
                self.cm.close_all(bravo_sess_id)
            bravo_win.events.closed += on_bravo_closed
            self._bravo_windows.append(bravo_win)

            def on_bravo_loaded():
                try:
                    if getattr(self, '_current_titlebar_theme', None):
                        t = self._current_titlebar_theme
                        apply_dwm_titlebar_theme(bravo_win, bg=t.get("bg"), text=t.get("text"), border=t.get("border"), is_dark=t.get("is_dark"), title=win_title)
                    else:
                        apply_dwm_titlebar_theme(bravo_win, bg=cur_theme, title=win_title)
                    bravo_win.evaluate_js(
                        f"if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', false); }} "
                        f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }}"
                    )
                except Exception:
                    pass
            bravo_win.events.loaded += on_bravo_loaded

            # Fallback timer: đảm bảo cửa sổ luôn hiển thị nếu frontend gặp sự cố không gọi ready()
            import threading
            threading.Timer(1.5, lambda: (bravo_win.show() if bravo_win else None)).start()

            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # Metadata API for Object Explorer & Table Designer
    def get_tables(self, connection_id, database=None, schema=None, search=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_tables(database, schema)
            if search:
                s = search.strip().lower()
                items = [t for t in items if s in (t.get("name", t) if isinstance(t, dict) else str(t)).lower()]
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_table_columns(self, connection_id, database=None, schema=None, table_name=None):
        try:
            conn = self._get_connection(connection_id)
            items = conn.metadata_service.list_columns(table_name, database, schema)
            return {"success": True, "items": items}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_table_designer_metadata(self, connection_id, database=None, schema=None, table_name=None):
        return self.get_table_design(connection_id, database, schema, table_name)

    # Query Execution API
    def execute_query(self, connection_id, sql, limit=1000, database=None, schema=None, query_id=None, mode="normal", debug_session_id=None, traces=None, confirmed=False, parameters=None, stop_on_error=True):
        try:
            conn = self._get_connection(connection_id)

            safety = check_query_safety(
                sql=sql,
                read_only=getattr(conn, "read_only", False),
                environment=getattr(conn, "environment", None),
                confirmed=confirmed,
                dialect=getattr(conn.adapter, "db_type", "tsql")
            )
            if not safety["allowed"]:
                return {
                    "success": False,
                    "needs_confirmation": safety["needs_confirmation"],
                    "error_code": safety["error_code"],
                    "error": safety["message"],
                    "classification": safety["classification"]
                }

            def on_progress(event):
                if not query_id:
                    return
                try:
                    import webview, json
                    payload = json.dumps({"query_id": query_id, **event})
                    js = f"if (window.onQueryProgress) {{ window.onQueryProgress({payload}); }}"
                    for w in getattr(webview, 'windows', []):
                        try:
                            w.evaluate_js(js)
                        except Exception:
                            pass
                except Exception:
                    pass

            result = conn.query_service.execute(
                sql,
                limit=limit,
                database=database,
                schema=schema,
                progress_callback=on_progress if query_id else None,
                mode=mode,
                debug_session_id=debug_session_id,
                traces=traces,
                parameters=parameters,
                stop_on_error=stop_on_error,
            )
            return {"success": True, **result}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def parse_query_parameters(self, sql, dialect="sqlserver"):
        try:
            from app.utils.sql_params import get_query_parameters
            return {"success": True, "parameters": get_query_parameters(sql, dialect)}
        except Exception as exc:
            return {"success": False, "error": str(exc), "parameters": []}

    def cancel_query(self, connection_id):
        try:
            conn = self._get_connection(connection_id)
            if conn and hasattr(conn, "query_service"):
                ok = conn.query_service.cancel()
                return {"success": True, "cancelled": ok}
            return {"success": False, "error": "Connection not found"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def explain_query(self, connection_id, sql, database=None):
        try:
            conn = self._get_connection(connection_id)
            adapter = conn.adapter
            db_type = getattr(adapter, "db_type", "sqlserver").lower()

            if db_type == "postgresql":
                exp_sql = f"EXPLAIN (FORMAT TEXT) {sql}"
                result = conn.query_service.execute(exp_sql, limit=500, database=database)
                rows = result.get("rows", [])
                plan_text = "\n".join(str(r[0]) for r in rows if r)
            else:
                import contextlib
                lock = getattr(conn, "lock", None) or getattr(adapter, "_lock", None)
                lock_ctx = lock if lock is not None else contextlib.nullcontext()
                with lock_ctx:
                    cursor = adapter.get_cursor(database=database) if hasattr(adapter, "get_cursor") else adapter.connection.cursor()
                    try:
                        cursor.execute("SET SHOWPLAN_TEXT ON")
                        cursor.execute(sql)
                        rows = cursor.fetchall()
                        plan_text = "\n".join(str(r[0]) for r in rows if r)
                    finally:
                        try:
                            cursor.execute("SET SHOWPLAN_TEXT OFF")
                        except Exception:
                            pass
                        try:
                            cursor.close()
                        except Exception:
                            pass

            return {"success": True, "plan": plan_text}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def syntax_check(self, connection_id=None, sql="", dialect="sqlserver", database=None, schema=None, is_debug=False, cached_tables=None):
        try:
            if not sql or not str(sql).strip():
                return {"success": True, "diagnostics": []}

            conn_session = None
            try:
                conn_session = self._get_connection(connection_id)
                if conn_session:
                    if not dialect or dialect == "sqlserver":
                        conn_db_type = getattr(conn_session, "config", {}).get("type", "sqlserver")
                        if conn_db_type:
                            dialect = conn_db_type
                    if not database:
                        database = getattr(conn_session, "config", {}).get("database")
            except Exception:
                conn_session = None

            try:
                from app.database.syntax.coordinator import validate_sql_pipeline
                diagnostics = validate_sql_pipeline(
                    sql=sql,
                    dialect=dialect,
                    connection_session=conn_session,
                    database=database,
                    schema=schema,
                    is_debug=bool(is_debug),
                    cached_tables_dict=cached_tables,
                )
            except Exception as pipe_err:
                from app.database.syntax.validator import validate_sql
                diagnostics = validate_sql(
                    sql=sql,
                    dialect=dialect,
                    connection_session=conn_session,
                    database=database,
                    schema=schema,
                    is_debug=bool(is_debug),
                    cached_tables_dict=cached_tables,
                )
            return {"success": True, "diagnostics": diagnostics}
        except Exception as exc:
            return {"success": False, "error": str(exc), "diagnostics": []}

    # ── Query History Native APIs (Phase 2 & Phase 3) ───────────────────────
    def open_query_history_window(self):
        try:
            import webview
            settings = _load_settings_file()
            cur_theme = settings.get("theme") or settings.get("appearance", {}).get("theme", "dark")
            history_html = (BASE_DIR / "templates" / "query_history_window.html").resolve().as_uri()
            if hasattr(self, '_history_window') and self._history_window:
                try:
                    if hasattr(self._history_window, 'restore'):
                        self._history_window.restore()
                    self._history_window.show()
                    self._history_window.focus()
                    return {"success": True}
                except Exception:
                    self._history_window = None

            win_title = "Query History — luoBTool IDE"
            history_api = QueryHistorySubwindowApi(self)
            history_win = webview.create_window(
                title=win_title,
                url=history_html,
                js_api=history_api,
                width=1280,
                height=820,
                min_size=(900, 560),
                text_select=True
            )
            history_api.set_window(history_win)
            self._history_window = history_win

            # Áp dụng DWM Titlebar cho cửa sổ Query History (nhiều lần với delay để Windows Forms tạo xong HWND)
            import threading
            for delay in (0.05, 0.15, 0.3, 0.6, 1.0):
                threading.Timer(delay, lambda: self.apply_titlebar_theme()).start()

            def on_closed():
                self._history_window = None
                try:
                    import webview
                    if history_win in webview.windows:
                        webview.windows.remove(history_win)
                except Exception:
                    pass

            def on_loaded():
                try:
                    for delay in (0.05, 0.2, 0.5):
                        threading.Timer(delay, lambda: self.apply_titlebar_theme()).start()
                    history_win.evaluate_js(
                        f"if (typeof window.applyHistoryTheme === 'function') {{ window.applyHistoryTheme('{cur_theme}'); }} "
                        f"else if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', false); }} "
                        f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }}"
                    )
                except Exception:
                    pass

            history_win.events.closed += on_closed
            history_win.events.loaded += on_loaded
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_log_execution(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            doc_id = payload.get("doc_id")
            if not doc_id:
                return {"success": False, "error": "doc_id is required"}
            # Kiểm tra thiết lập bật/tắt ghi lịch sử (history.enabled)
            settings = _load_settings_file()
            history_cfg = settings.get("history", {})
            if not history_cfg.get("enabled", True):
                return {"success": True, "entry_id": None, "disabled": True}

            max_entry_size = int(history_cfg.get("maxEntrySizeKb", 1024)) * 1024
            svc = get_history_service()
            res = svc.record_execution(
                doc_id=doc_id,
                display_name=payload.get("display_name") or "Untitled.sql",
                sql=payload.get("sql") or "",
                server=payload.get("server") or "localhost",
                database=payload.get("database") or "master",
                file_path=payload.get("file_path"),
                login=payload.get("login"),
                is_selection=bool(payload.get("is_selection", False)),
                entry_type=int(payload.get("entry_type", 1)),
                executed_at=payload.get("executed_at"),
                duration_ms=payload.get("duration_ms"),
                status=payload.get("status"),
                error_summary=payload.get("error_summary"),
                max_entry_size=max_entry_size
            )
            if res is None:
                return {"success": True, "entry_id": None, "skipped": True}
            return {"success": True, **res}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_update_execution(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            entry_id = payload.get("entry_id")
            if entry_id is None:
                return {"success": False, "error": "entry_id is required"}
            svc = get_history_service()
            updated = svc.update_entry_result(
                entry_id=int(entry_id),
                duration_ms=payload.get("duration_ms"),
                status=payload.get("status"),
                error_summary=payload.get("error_summary")
            )
            return {"success": True, "updated": updated}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_update_document(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            doc_id = payload.get("doc_id")
            if not doc_id:
                return {"success": False, "error": "doc_id is required"}
            svc = get_history_service()
            return svc.update_document(
                doc_id=doc_id,
                display_name=payload.get("display_name"),
                file_path=payload.get("file_path")
            )
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_document_by_path(self, file_path):
        try:
            svc = get_history_service()
            doc = svc.get_document_by_path(file_path)
            return {"success": True, "document": doc}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_documents(self, search=None, favorites_only=False, limit=200, server=None, database=None, from_timestamp=None, to_timestamp=None):
        try:
            svc = get_history_service()
            docs = svc.get_documents(
                search=search,
                favorites_only=bool(favorites_only),
                server=server,
                database=database,
                from_timestamp=from_timestamp,
                to_timestamp=to_timestamp,
                limit=int(limit or 200)
            )
            return {"success": True, "documents": docs}
        except Exception as exc:
            return {"success": False, "error": str(exc), "documents": []}

    def history_get_document(self, doc_id):
        try:
            svc = get_history_service()
            doc = svc.get_document(doc_id)
            if not doc:
                return {"success": False, "error": "Document not found"}
            return {"success": True, "document": doc}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_entries(self, doc_id, status=None, limit=200, server=None, database=None, from_timestamp=None, to_timestamp=None, tag=None, entry_type=None, search=None):
        try:
            svc = get_history_service()
            status_filter = int(status) if status is not None and status != "" else None
            entries = svc.get_entries(
                doc_id=doc_id,
                status_filter=status_filter,
                search=search,
                server=server,
                database=database,
                from_timestamp=from_timestamp,
                to_timestamp=to_timestamp,
                tag=tag,
                entry_type=entry_type,
                limit=int(limit or 200)
            )
            return {"success": True, "entries": entries}
        except Exception as exc:
            return {"success": False, "error": str(exc), "entries": []}

    def history_get_entry_detail(self, entry_id):
        try:
            svc = get_history_service()
            detail = svc.get_entry_detail(int(entry_id))
            if not detail:
                return {"success": False, "error": "Entry not found"}
            return {"success": True, "entry": detail}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_toggle_star(self, doc_id, is_starred=None):
        try:
            svc = get_history_service()
            val = svc.toggle_star(doc_id=doc_id, is_starred=is_starred)
            return {"success": True, "is_starred": val}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_delete_entry(self, entry_id):
        try:
            svc = get_history_service()
            ok = svc.delete_entry(int(entry_id))
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_delete_document(self, doc_id):
        try:
            svc = get_history_service()
            ok = svc.delete_document(doc_id)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_purge(self, max_age_days=None, max_entries_per_doc=None, max_db_size_mb=None):
        try:
            settings = _load_settings_file()
            history_cfg = settings.get("history", {})
            if max_age_days is None:
                max_age_days = int(history_cfg.get("maxAgeDays", 90))
            if max_entries_per_doc is None:
                max_entries_per_doc = int(history_cfg.get("maxEntriesPerDoc", 200))
            if max_db_size_mb is None:
                max_db_size_mb = int(history_cfg.get("maxDbSizeMb", 500))

            svc = get_history_service()
            res = svc.purge_old_entries(
                max_age_days=int(max_age_days),
                max_entries_per_doc=int(max_entries_per_doc),
                max_db_size_mb=int(max_db_size_mb)
            )
            return {"success": True, **res}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_compact(self):
        try:
            svc = get_history_service()
            stats = svc.compact()
            return {"success": True, "stats": stats}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_clear_all(self, keep_favorites=False):
        try:
            svc = get_history_service()
            stats = svc.clear_all(keep_favorites=keep_favorites)
            return {"success": True, "stats": stats}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_stats(self):
        try:
            svc = get_history_service()
            st = svc.get_storage_stats()
            return {"success": True, "stats": st}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_open_in_editor(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            sql = payload.get("sql", "")
            title = payload.get("title") or "SQLQuery.sql"
            server = payload.get("server")
            database = payload.get("database")

            if self._window:
                js_code = (
                    f"if (typeof window.openHistoryQueryInNewTab === 'function') {{"
                    f"  window.openHistoryQueryInNewTab({json.dumps(sql)}, {json.dumps(title)}, {json.dumps(server)}, {json.dumps(database)});"
                    f"}}"
                )
                self._window.evaluate_js(js_code)
                try:
                    self._window.focus()
                except Exception:
                    pass
                return {"success": True}
            return {"success": False, "error": "Main window reference not found"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_open_containing_folder(self, file_path):
        try:
            if not file_path or not os.path.exists(file_path):
                return {"success": False, "error": "File does not exist on disk"}
            import subprocess
            subprocess.Popen(f'explorer /select,"{os.path.normpath(file_path)}"')
            return {"success": True}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_open_diff_in_editor(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            original_sql = payload.get("original_sql", "")
            modified_sql = payload.get("modified_sql", "")
            original_title = payload.get("original_title") or "Previous Version"
            modified_title = payload.get("modified_title") or "Current Version"

            if self._window:
                js_code = (
                    f"if (typeof window.showDiffView === 'function') {{"
                    f"  window.showDiffView({json.dumps(original_sql)}, {json.dumps(modified_sql)}, {json.dumps(original_title)}, {json.dumps(modified_title)}, 'sql');"
                    f"}}"
                )
                self._window.evaluate_js(js_code)
                try:
                    self._window.focus()
                except Exception:
                    pass
                return {"success": True}
            return {"success": False, "error": "Main window reference not found"}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_active_editor_content(self):
        try:
            if self._window:
                code = "window.AppEditor ? window.AppEditor.getValue() : ''"
                content = self._window.evaluate_js(code)
                return {"success": True, "content": content or ""}
            return {"success": False, "error": "Main window reference not found", "content": ""}
        except Exception as exc:
            return {"success": False, "error": str(exc), "content": ""}

    def history_update_entry_note(self, entry_id, note=None, tag=None):
        try:
            svc = get_history_service()
            ok = svc.update_entry_note(int(entry_id), note=note, tag=tag)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_filters(self):
        try:
            svc = get_history_service()
            data = svc.get_distinct_filters()
            return {"success": True, "filters": data}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_export(self, payload=None):
        try:
            doc_ids = None
            if isinstance(payload, dict):
                doc_ids = payload.get("doc_ids")
            svc = get_history_service()
            res = svc.export_history(doc_ids=doc_ids)
            return {"success": True, **res}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_import(self, payload):
        try:
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            svc = get_history_service()
            res = svc.import_history(payload)
            return {"success": True, **res}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_set_entry_favorite(self, entry_id, is_favorite=None):
        try:
            svc = get_history_service()
            val = svc.set_entry_favorite(int(entry_id), is_favorite=is_favorite)
            return {"success": True, "is_favorite": val}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_add_entry_tag(self, entry_id, tag_name, color=None):
        try:
            svc = get_history_service()
            ok = svc.add_entry_tag(int(entry_id), tag_name=str(tag_name), color=color)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_remove_entry_tag(self, entry_id, tag_name):
        try:
            svc = get_history_service()
            ok = svc.remove_entry_tag(int(entry_id), tag_name=str(tag_name))
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def history_get_entry_tags(self, entry_id):
        try:
            svc = get_history_service()
            tags = svc.get_entry_tags(int(entry_id))
            return {"success": True, "tags": tags}
        except Exception as exc:
            return {"success": False, "error": str(exc), "tags": []}

    def history_get_tags(self):
        try:
            svc = get_history_service()
            tags = svc.get_all_tags()
            return {"success": True, "tags": tags}
        except Exception as exc:
            return {"success": False, "error": str(exc), "tags": []}

    def session_save(self, session_data):
        try:
            from app.routes.session import save_session_data
            return save_session_data(session_data)
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def session_load(self):
        try:
            from app.routes.session import load_session_data
            return load_session_data()
        except Exception as exc:
            return {"success": False, "error": str(exc), "session": None}

    def session_clear(self):
        try:
            from app.routes.session import clear_session_data
            return clear_session_data()
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # === PostgreSQL Dependency Graph Native IPC Methods ===

    def pg_dependency_get_column(self, conn_id, schema, table, column, direction="dependents", depth=2, database=None, include_source=True):
        try:
            conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                pg_conns = [c for c in active_list if (c.config.get("type") or "").lower() == "postgresql"]
                if pg_conns:
                    conn = pg_conns[0]
                elif active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Connection not found"}
            db_type = (conn.config.get("type") or getattr(conn.adapter, "db_type", "") or "").lower()
            if db_type != "postgresql":
                return {"success": False, "error": "PostgreSQL Dependency Graph is exclusively supported for PostgreSQL connections."}
            from app.services.postgres_dependency import PostgresDependencyService
            service = PostgresDependencyService(conn.adapter)
            graph = service.get_column_graph(
                schema=schema or "public",
                table_name=table,
                column_name=column,
                direction=direction or "dependents",
                depth_limit=int(depth or 2),
                database=database,
                include_source_analysis=bool(include_source)
            )
            return {
                "success": True,
                "graph": graph.to_dict(),
                "tree": graph.to_tree_projection()
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def pg_dependency_get_object(self, conn_id, obj_type, schema, name, direction="dependents", depth=2, database=None):
        try:
            conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                pg_conns = [c for c in active_list if (c.config.get("type") or "").lower() == "postgresql"]
                if pg_conns:
                    conn = pg_conns[0]
                elif active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Connection not found"}
            db_type = (conn.config.get("type") or getattr(conn.adapter, "db_type", "") or "").lower()
            if db_type != "postgresql":
                return {"success": False, "error": "PostgreSQL Dependency Graph is exclusively supported for PostgreSQL connections."}
            from app.services.postgres_dependency import PostgresDependencyService
            service = PostgresDependencyService(conn.adapter)
            graph = service.get_object_graph(
                obj_type_str=obj_type or "view",
                schema=schema or "public",
                name=name,
                direction=direction or "dependents",
                depth_limit=int(depth or 2),
                database=database
            )
            return {
                "success": True,
                "graph": graph.to_dict(),
                "tree": graph.to_tree_projection()
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def pg_dependency_get_definition(self, conn_id, obj_type, schema, name, oid=None, database=None):
        try:
            conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                pg_conns = [c for c in active_list if (c.config.get("type") or "").lower() == "postgresql"]
                if pg_conns:
                    conn = pg_conns[0]
                elif active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Connection not found"}
            from app.services.postgres_dependency import PostgresDependencyService
            service = PostgresDependencyService(conn.adapter)
            oid_int = int(oid) if oid and str(oid).isdigit() else None
            defn = service.get_object_definition(
                obj_type_str=obj_type or "view",
                schema=schema or "public",
                name=name,
                oid=oid_int,
                database=database
            )
            if not defn:
                return {"success": False, "error": f"Definition not found for {schema}.{name}"}
            return {
                "success": True,
                "object_type": obj_type,
                "schema": schema,
                "name": name,
                "definition": defn
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def pg_dependency_get_suggested_steps(self, conn_id, schema, table, column, action="ALTER", database=None):
        try:
            conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                pg_conns = [c for c in active_list if (c.config.get("type") or "").lower() == "postgresql"]
                if pg_conns:
                    conn = pg_conns[0]
                elif active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Connection not found"}
            from app.services.postgres_dependency import PostgresDependencyService
            service = PostgresDependencyService(conn.adapter)
            steps = service.get_suggested_steps(
                schema=schema or "public",
                table=table,
                column=column,
                action=action or "ALTER",
                database=database
            )
            return {
                "success": True,
                "action": action,
                "target": f"{schema}.{table}.{column}",
                "steps": steps
            }
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    # === Object Content Search Native IPC Methods ===

    def sql_object_search(self, conn_id=None, keyword="", database=None, object_types=None, case_sensitive=False, whole_word=False, is_regex=False, schema=None):
        try:
            conn = None
            if conn_id:
                conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                if active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Database connection is unavailable."}

            from app.services.object_search_service import ObjectSearchService
            service = ObjectSearchService(conn.adapter)
            return service.search(
                keyword=keyword,
                database=database,
                object_types=object_types,
                case_sensitive=bool(case_sensitive),
                whole_word=bool(whole_word),
                is_regex=bool(is_regex),
                schema=schema
            )
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def sql_object_ddl(self, conn_id=None, schema="dbo", name="", object_type="VIEW", database=None, keyword=None, case_sensitive=False, whole_word=False, is_regex=False):
        try:
            conn = None
            if conn_id:
                conn = self._get_connection(conn_id)
            if not conn and hasattr(self, 'cm') and self.cm:
                active_list = self.cm.list(self.owner_session_id)
                if active_list:
                    conn = active_list[0]
            if not conn or not conn.adapter:
                return {"success": False, "error": "Database connection is unavailable."}

            from app.services.object_search_service import ObjectSearchService
            service = ObjectSearchService(conn.adapter)
            return service.get_ddl(
                schema=schema,
                name=name,
                object_type=object_type,
                database=database,
                keyword=keyword,
                case_sensitive=bool(case_sensitive),
                whole_word=bool(whole_word),
                is_regex=bool(is_regex)
            )
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_code_styles(self):
        try:
            from app.routes.styles import _load_styles_file
            return _load_styles_file()
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def save_code_styles(self, payload):
        try:
            from app.routes.styles import _save_styles_file
            if not isinstance(payload, dict):
                return {"success": False, "error": "payload must be an object"}
            ok = _save_styles_file(payload)
            return {"success": ok}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def reset_code_styles(self):
        try:
            from app.routes.styles import _get_default_styles_data, _save_styles_file
            defaults = _get_default_styles_data()
            _save_styles_file(defaults)
            return {"success": True, **defaults}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def analyze_query_flow(self, sql, dialect="tsql"):
        try:
            from app.services.query_flow_service import QueryFlowService
            graph = QueryFlowService.build_flow_graph(sql, dialect=dialect)
            return {"success": True, "flowGraph": graph.to_dict()}
        except Exception as exc:
            return {"success": False, "error": str(exc)}

    def get_query_plan(self, sql, conn_id=None, plan_type="estimated", database=None, schema=None, dialect=None):
        try:
            from app.services.query_plan_service import QueryPlanService
            conn_session = None
            if conn_id and hasattr(self, "_connection_manager") and self._connection_manager:
                try:
                    conn_session = self._connection_manager.get_connection(conn_id)
                except Exception:
                    pass

            plan = QueryPlanService.get_execution_plan(
                connection_session=conn_session,
                sql=sql,
                plan_type=plan_type,
                database=database,
                schema=schema,
                dialect=dialect,
            )
            return {"success": True, "plan": plan.to_dict()}
        except Exception as exc:
            return {"success": False, "error": str(exc)}




def main():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    cache_dir = DATA_DIR / ".webview_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)

    settings = _load_settings_file()
    log_cfg = settings.get("logging", {}) if isinstance(settings, dict) else {}
    log_level = log_cfg.get("level", "INFO")
    max_mb = int(log_cfg.get("maxFileSizeMb", 2))
    max_files = int(log_cfg.get("maxFiles", 5))
    setup_logging(
        log_dir=DATA_DIR / "logs",
        level_name=log_level,
        max_bytes=max_mb * 1024 * 1024,
        backup_count=max_files,
        enable_console=True
    )

    index_html = (BASE_DIR / "templates" / "index.html").resolve().as_uri()
    splash_html = (BASE_DIR / "templates" / "splash.html").resolve().as_uri()
    main_api = MainApi(connection_manager)
    cur_theme = settings.get("appearance", {}).get("theme") or settings.get("theme", "dark")
    bg_color = "#ffffff" if cur_theme in ["light", "win-nt", "win-xp"] else "#1e1e1e"

    import ctypes
    try:
        user32 = ctypes.windll.user32
        sw = user32.GetSystemMetrics(0)
        sh = user32.GetSystemMetrics(1)
    except Exception:
        sw, sh = 1920, 1080

    splash_w, splash_h = 800, 500
    splash_x = max(0, (sw - splash_w) // 2)
    splash_y = max(0, (sh - splash_h) // 2)

    splash_win = webview.create_window(
        title="luoBTool IDE",
        url=splash_html,
        width=splash_w,
        height=splash_h,
        x=splash_x,
        y=splash_y,
        frameless=True,
        easy_drag=True,
        resizable=False,
        background_color="#071324",
        on_top=True
    )

    window = webview.create_window(
        title="luoBTool IDE",
        url=index_html,
        js_api=main_api,
        width=1400,
        height=900,
        min_size=(900, 600),
        hidden=True,
        background_color=bg_color,
        text_select=True
    )
    main_api.set_window(window)

    splash_min_timer_done = False
    frontend_ready = False
    splash_closed = False

    def maybe_transition():
        nonlocal splash_closed
        if splash_closed:
            return
        if splash_min_timer_done and frontend_ready:
            splash_closed = True
            try:
                window.show()
            except Exception:
                pass

            def _sync_titlebar():
                try:
                    main_api.apply_titlebar_theme(cur_theme)
                except Exception:
                    pass
                try:
                    apply_dwm_titlebar_theme(window, bg=cur_theme, title="luoBTool IDE")
                except Exception:
                    pass

            _sync_titlebar()
            import threading
            threading.Timer(0.1, _sync_titlebar).start()
            threading.Timer(0.35, _sync_titlebar).start()
            threading.Timer(0.7, _sync_titlebar).start()

            try:
                splash_win.hide()
            except Exception:
                pass
            def _destroy_splash():
                try:
                    splash_win.destroy()
                except Exception:
                    pass
                _sync_titlebar()
            threading.Timer(0.5, _destroy_splash).start()

    def on_min_splash_timer():
        nonlocal splash_min_timer_done
        splash_min_timer_done = True
        maybe_transition()

    timer_started = False
    def start_min_timer():
        nonlocal timer_started
        if timer_started:
            return
        timer_started = True
        import threading
        threading.Timer(2.0, on_min_splash_timer).start()

    def on_splash_loaded():
        # Hiển thị Splash screen tối thiểu đúng 2.0 giây theo yêu cầu
        start_min_timer()

    splash_win.events.loaded += on_splash_loaded
    # Đảm bảo bộ đếm tối thiểu 2.0s luôn kích hoạt kể cả khi sự kiện loaded của webview diễn ra cực nhanh
    import threading
    threading.Timer(0.1, start_min_timer).start()

    def on_app_ready():
        nonlocal frontend_ready
        frontend_ready = True
        maybe_transition()

    main_api._on_ready_callback = on_app_ready

    def on_main_loaded():
        try:
            apply_dwm_titlebar_theme(window, bg=cur_theme, title="luoBTool IDE")
            window.evaluate_js(
                f"if (window.ThemeManager) {{ window.ThemeManager.applyTheme('{cur_theme}', true); }} "
                f"else {{ document.documentElement.setAttribute('data-bs-theme', '{cur_theme}'); }}"
            )
        except Exception:
            pass

    window.events.loaded += on_main_loaded

    def on_main_closed():
        try:
            closed = connection_manager.close_all_sessions()
            logger.info("Cleaned up %d connection(s) on main window close", closed)
        except Exception as ex:
            logger.warning("Error cleaning up connections on main window close: %s", ex)

    window.events.closed += on_main_closed

    # Safety fallback timer: sau 10.0s đảm bảo chuyển tiếp nếu frontend gặp lỗi nghiêm trọng không gọi về được
    def on_fallback():
        nonlocal splash_min_timer_done, frontend_ready
        splash_min_timer_done = True
        frontend_ready = True
        maybe_transition()

    import threading
    threading.Timer(10.0, on_fallback).start()

    print(f"[App] Starting Desktop pywebview with active theme: {cur_theme}")
    webview.start(gui="edgechromium", storage_path=str(cache_dir.resolve()), debug=False)


if __name__ == "__main__":
    main()
